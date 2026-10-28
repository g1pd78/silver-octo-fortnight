"""
Прогноз ЭФФЕКТИВНОСТИ сборов (факт / KPI-план) — на каждую дату месяца и на его конец, по итогу и по продуктам.

ГЛАВНАЯ МЕТРИКА
  Эффективность на дату = накопленный факт / накопленный план с начала месяца (на конец месяца — факт / план месяца).
  Прогноз эффективности = накопленный прогноз / накопленный план. Цель — разница с фактом не больше 1 п.п.
  Бэктест, выбор метода и отчёт считаются по этой разнице: на конец месяца (в зависимости от того, сколько дней
  месяца уже известно) и на каждую дату. Пока KPI месяца не загружен (он приходит в первые дни), эффективность
  прогноза считается к спрогнозированному плану — режим «план ещё не загружен».

КАК УСТРОЕН ПРОГНОЗ
  1. Календарные признаки по производственному календарю РФ 2024–2027: выходные, праздники, каникулы,
     дни зарплаты, рабочий день по счёту, длинные выходные до и после дня.
  2. Прогнозируется факт в деньгах на каждый день, эффективность = прогноз / план. KPI — цель бизнеса, равномерно
     разложенная по дням: ни его уровень, ни раскладка не говорят, сколько и когда заплатят клиенты. Поэтому кроме
     методов «к плану» (как раньше) есть методы, которым KPI не нужен: от «к оплате» и от тренда факта.
  3. «К оплате на дату» (DEBT_COL), известное на весь месяц вперёд, — главный источник точности: деньги приходят
     вокруг дат платежа, и сколько придёт, зависит от того, сколько к оплате. Модели от «к оплате»: ожидание оплат
     дня по датам платежа E(t) = Σ w_l · «к оплате»(t + l) — веса (кто платит заранее, в срок, позже) оцениваются по
     истории — × поправка бустинга (календарь, праздники, дни зарплаты, recovery прошлых месяцев). Ожидание линейно
     по «к оплате»: всплеск первого рабочего дня после длинных праздников, больше любого в истории, масштабируется,
     а не обрезается, как в деревьях.
  4. Внутри месяца остаток прогноза поправляется по факту на сегодня. Доверие к факту растёт с долей месяца, которая
     уже должна была прийти: α = α_max · c / (c + τ), c — ожидаемая по прогнозу доля месяца по сегодня. В первые
     шумные дни поправка слабая, к концу месяца — сильная.
  5. Честный бэктест: для каждого месяца метод, веса смеси и параметры выбираются только по предыдущим месяцам —
     так же, как прогноз работает в жизни. «Авто» в отчёте — ожидаемая точность рабочего прогноза.
  6. Продукты (SEGMENT_COL, run_segments): прогноз эффективности по каждому продукту и по итогу. Итог — модель на сумме
     или сумма прогнозов продуктов: что точнее в честном бэктесте, то и идёт в прогноз итога.

МЕТОДЫ (прогноз месяца без его факта; дальше — поправка по факту на сегодня)
  Референсные месяцы          — текущий подход: средняя кривая выполнения 3 прошлых месяцев по числу месяца.
  Реф. месяцы по раб. дням    — то же, но дни выровнены по номеру рабочего дня от начала и от конца месяца.
  Бустинг к плану             — бустинг учит долю плана месяца на каждый день: календарь, раскладка плана и
                                «к оплате», выполнение прошлых месяцев, «к оплате» к плану.
  Календарь × уровень к плану — тот же бустинг без признаков уровня × выполнение плана, сглаженное по прошлым месяцам.
  Календарь × тренд факта     — календарная модель × тренд факта прошлых месяцев (KPI не нужен).
  Бустинг от «к оплате»       — ожидание оплат по датам платежа × поправка бустинга (нужен DEBT_COL).
  Календарь × recovery        — то же без признаков уровня × recovery, сглаженный по прошлым месяцам.
  Смесь                       — взвешенная сумма лучших методов.
  Авто                        — каждый месяц лучший по предыдущим месяцам вариант (метод или смесь с параметрами).

ВЫХОД
  forecast_<месяц>.csv — по каждому дню: план, факт, forecast — прогноз на каждый день без факта этого месяца
  (для текущего месяца — сделанный на 1-е число, по нему проверка по факту), fact_plus_forecast — факт, где он есть,
  + прогноз остатка с поправкой по факту на сегодня; накопленные суммы; эффективность на дату (eff_*_pct) и разница
  прогноза с фактом (eff_diff_pp); recovery rate, если задан DEBT_COL. В сводке — эффективность на конец месяца
  с 80% интервалом из честного бэктеста. backtest_errors.csv — честный бэктест по месяцам и дням.
  effectiveness.xlsx — всё в Excel, в формате треугольника: строки — месяцы, столбцы — числа месяца 1–31, в ячейке —
  эффективность на эту дату. Листы: «Эффективность» (факт, где его ещё нет — прогноз), «Прогноз с 1-го числа»,
  «Разница с фактом» (по дням, п.п.), «Итог месяца по мере факта» (прогноз на конец месяца, когда известно k дней),
  сводки и прогноз по дням.

ВХОД
  DataFrame с колонками дата / факт / план [/ «к оплате»] [/ продукт] — значения ЗА ДЕНЬ. Будущие даты факта не нужны.
  План месяца, пока он не загружен, оставьте пустым. «К оплате» выгружайте на все дни прогнозного месяца.
  ВАЖНО про историю «к оплате»: в бэктесте «к оплате» прошлого месяца должно быть таким, каким его знали на 1-е число
  этого месяца (по графикам). Если в истории лежат значения «на дату», уже изменённые оплатами внутри месяца (клиент
  заплатил раньше — и «к оплате» на дату платежа уменьшилось), бэктест покажет точность лучше, чем будет в жизни.
  Проверка: сравните «к оплате» на даты одного месяца из истории и из выгрузки, сделанной на 1-е число этого месяца.

Зависимости: pandas, numpy, scikit-learn, matplotlib. lightgbm — по желанию, используется, если установлен.
Запуск: заполните load_data() и выполните файл, или в Jupyter:
  import collections_forecast as cf
  cf.DEBT_COL = "debt"                     # колонка «к оплате»
  res = cf.run(df)                         # общий итог
  res = cf.run_segments(df, "product")     # каждый продукт и итог
Настройки — в блоке НАСТРОЙКИ ниже. Прогон итога 2–4 минуты с lightgbm (без него дольше; N_SEEDS = 1 — около минуты,
для черновых прогонов); run_segments — столько же на каждый продукт и на итог.
"""

import datetime
import itertools
import warnings

import numpy as np
import pandas as pd

# ============================================================================
# НАСТРОЙКИ
# ============================================================================
DATE_COL, FACT_COL, PLAN_COL = "date", "fact", "plan"
DEBT_COL = None                # «к оплате на дату» (сколько клиенты должны внести на дату, включая неоплаченное ранее),
                               # например "debt"; None — без него
DEBT_KNOWN_AHEAD = True        # True — «к оплате» на все дни месяца известно на 1-е число (выгружайте его на весь
                               # прогнозный месяц): модели от «к оплате» по дням. False — известно только по сегодня:
                               # только месячные признаки (recovery прошлых месяцев, «к оплате» к плану)
SEGMENT_COL = None             # колонка продукта/сегмента для run_segments(), например "product"
CUMULATIVE_INPUT = False       # True, если факт и план в выгрузке накопленные с начала месяца
LAST_FACT_DATE = None          # последняя дата с ПОЛНЫМ фактом. None = последняя дата, где факт не пустой.
                               # Если в выгрузке есть неполный «сегодня», укажите вчерашнюю дату.
TARGET_MONTHS = None           # за какие месяцы выдать полный прогноз по дням, например ["2026-10", "2026-11"];
                               # None = текущий месяц и HORIZON_MONTHS следующих
HORIZON_MONTHS = 1             # сколько месяцев после текущего прогнозировать, если TARGET_MONTHS = None
RR_DENOM = "sum"               # recovery rate на дату: "sum" — Σ факт / Σ «к оплате» за те же дни;
                               # "last" — Σ факт / долг на эту дату (если долг — остаток портфеля на дату)
BACKTEST_FROM = "2025-01"      # с какого месяца считать бэктест (более ранние месяцы — только обучение)
HONEST_WARMUP = 6              # честный бэктест: первые N месяцев бэктеста — только для выбора метода, оценка — дальше
PAYDAYS = (5, 10, 15, 20, 25)  # типичные дни зарплаты/аванса клиентов; выходной -> предыдущий рабочий день
N_REF_MONTHS = 3               # сколько предыдущих месяцев брать в методе референсных месяцев
RECENCY_HALFLIFE = 12          # вес наблюдения падает вдвое каждые N месяцев давности (None — без весов)
EXCLUDE_DATES = []             # аномальные даты, которые не брать в обучение, например ["2025-03-14"]
REPORT_KNOWN_DAYS = (0, 1, 3, 5, 10, 15, 20, 25)  # «сколько дней месяца уже известно» в отчёте
NOWCAST_MAX = (0.25, 0.5, 0.75, 1.0)              # сила поправки остатка по факту месяца (α_max)
NOWCAST_TAU = (0, 0.02, 0.05, 0.1, 0.2, 0.5, 1.0)  # доверие к факту c / (c + τ): чем больше τ, тем позже доверяем
LEVEL_HALFLIVES = (1, 2, 3, 6, 12, np.inf)   # «× уровень», «× recovery»: вес прошлого месяца падает вдвое каждые N мес.
LEVEL_PRIOR_WEIGHTS = (0, 1, 3, 10, np.inf)  # тяга уровня к среднему модели (inf = не подстраивать)
MIX_TOP = 3                    # смесь — из стольких лучших методов
DUE_LAGS = range(-7, 8)        # модели от «к оплате»: оплаты за 7 дней до даты платежа (+) … через 7 дней после (−)
PROFILE_MONTHS = 6             # по скольким последним месяцам раскладывать прогноз плана/«к оплате» по дням
TREND_HALFLIVES = (0, 1, 3, 6, 12, np.inf)  # темп роста: 0 = без роста, иначе полураспад весов, мес.
OUT_DIR = "."
EXCEL = True                   # выгружать результаты в Excel с треугольниками эффективности (нужен openpyxl)
RANDOM_STATE = 42
N_SEEDS = 5                    # сколько бустингов с разными seed усреднять (1 — быстрее, но прогноз «дрожит»)

# ============================================================================
# ПРОИЗВОДСТВЕННЫЙ КАЛЕНДАРЬ РФ
# Нерабочие дни (праздники и перенесённые выходные). Обычные сб/вс считаются автоматически.
# 2024–2026: КонсультантПлюс / Главкнига; 2027: постановление № 1187 от 17.09.2026 (по данным БУХ.1С).
# При появлении календаря на следующий год просто допишите его сюда.
# ============================================================================
RU_NONWORKING = {
    2024: ["01-01", "01-02", "01-03", "01-04", "01-05", "01-06", "01-07", "01-08",
           "02-23", "03-08", "04-29", "04-30", "05-01", "05-09", "05-10",
           "06-12", "11-04", "12-30", "12-31"],
    2025: ["01-01", "01-02", "01-03", "01-04", "01-05", "01-06", "01-07", "01-08",
           "05-01", "05-02", "05-08", "05-09", "06-12", "06-13",
           "11-03", "11-04", "12-31"],
    2026: ["01-01", "01-02", "01-03", "01-04", "01-05", "01-06", "01-07", "01-08", "01-09",
           "02-23", "03-09", "05-01", "05-11", "06-12", "11-04", "12-31"],
    2027: ["01-01", "01-02", "01-03", "01-04", "01-05", "01-06", "01-07", "01-08",
           "02-22", "02-23", "03-08", "05-03", "05-10", "06-14", "11-04", "11-05", "12-31"],
}
# Рабочие субботы (перенесённые рабочие дни)
RU_WORKING_WEEKENDS = ["2024-04-27", "2024-11-02", "2024-12-28", "2025-11-01", "2027-02-20"]


def build_calendar(start, end, paydays=PAYDAYS) -> pd.DataFrame:
    """Календарные признаки на каждый день диапазона (с запасом по краям)."""
    pad = pd.Timedelta(days=40)
    dates = pd.date_range(pd.Timestamp(start) - pad, pd.Timestamp(end) + pad, freq="D")
    missing = sorted(set(dates.year) - set(RU_NONWORKING))
    inside = [y for y in missing if pd.Timestamp(start).year <= y <= pd.Timestamp(end).year]
    if inside:
        warnings.warn(f"Нет производственного календаря на {inside}: праздники там не учтены, допишите RU_NONWORKING")

    nonwork = pd.to_datetime([f"{y}-{md}" for y, mds in RU_NONWORKING.items() for md in mds])
    working_weekends = pd.to_datetime(RU_WORKING_WEEKENDS)

    cal = pd.DataFrame({"date": dates})
    dow = cal["date"].dt.day_of_week
    is_nonwork = cal["date"].isin(nonwork)
    is_work = ((dow < 5) & ~is_nonwork) | cal["date"].isin(working_weekends)

    cal["dow"] = dow
    cal["day"] = cal["date"].dt.day
    cal["month"] = cal["date"].dt.month
    cal["dim"] = cal["date"].dt.days_in_month
    cal["days_to_eom"] = cal["dim"] - cal["day"]
    cal["is_workday"] = is_work.astype(int)
    cal["is_holiday"] = (is_nonwork & (dow < 5)).astype(int)  # будний день, но нерабочий

    ym = cal["date"].dt.to_period("M")
    cal["wd_num"] = cal.groupby(ym)["is_workday"].cumsum()
    cal["wd_total"] = cal.groupby(ym)["is_workday"].transform("sum")
    cal["wd_left"] = cal["wd_total"] - cal["wd_num"]
    cal["wd_frac"] = cal["wd_num"] / cal["wd_total"]

    # сколько нерабочих дней подряд было прямо перед днём / будет сразу после
    w = cal["is_workday"].to_numpy()
    n = len(w)
    before, after = np.zeros(n, int), np.zeros(n, int)
    run = 0
    for i in range(n):
        before[i] = run
        run = 0 if w[i] else run + 1
    run = 0
    for i in range(n - 1, -1, -1):
        after[i] = run
        run = 0 if w[i] else run + 1
    cal["nonwork_before"] = before
    cal["nonwork_after"] = after
    cal["first_wd_month"] = ((cal["is_workday"] == 1) & (cal["wd_num"] == 1)).astype(int)
    cal["last_wd_month"] = ((cal["is_workday"] == 1) & (cal["wd_left"] == 0)).astype(int)

    # новогодние каникулы: после последнего рабочего дня декабря и до первого рабочего дня января
    wd = cal.loc[cal["is_workday"] == 1, "date"]
    first_jan = wd[wd.dt.month == 1].groupby(wd.dt.year).min()
    last_dec = wd[wd.dt.month == 12].groupby(wd.dt.year).max()
    y = cal["date"].dt.year
    in_jan = (cal["month"] == 1) & (cal["date"] < y.map(first_jan))
    in_dec = (cal["month"] == 12) & (cal["date"] > y.map(last_dec))
    cal["ny_break"] = (in_jan | in_dec).astype(int)

    # дни зарплаты: если выпадают на нерабочий день — переносятся на предыдущий рабочий (ст. 136 ТК)
    wd_arr = wd.to_numpy()
    months = pd.period_range(dates.min(), dates.max(), freq="M")
    all_paydays = []
    for p in paydays:
        eff = []
        for m in months:
            target = np.datetime64(pd.Timestamp(m.year, m.month, min(p, m.days_in_month)))
            idx = np.searchsorted(wd_arr, target, side="right") - 1
            if idx >= 0:
                eff.append(wd_arr[idx])
        eff = pd.to_datetime(eff)
        cal[f"pay_{p}"] = cal["date"].isin(eff).astype(int)
        all_paydays.extend(eff)
    pays = np.sort(pd.to_datetime(all_paydays).unique().to_numpy())
    idx = np.searchsorted(pays, cal["date"].to_numpy(), side="right") - 1
    since = np.where(idx >= 0, (cal["date"].to_numpy() - pays[np.clip(idx, 0, None)]) / np.timedelta64(1, "D"), np.nan)
    cal["days_since_pay"] = since
    return cal


# ============================================================================
# ПОДГОТОВКА ДАННЫХ
# ============================================================================
def _cum_to_daily(s: pd.Series) -> pd.Series:
    last = s.last_valid_index()
    if last is None:
        return s
    f = s.loc[:last].ffill()
    return f.diff().fillna(f).reindex(s.index)


def target_months(cur):
    if TARGET_MONTHS:
        return sorted(pd.Period(m, "M") for m in TARGET_MONTHS)
    return [cur + i for i in range(HORIZON_MONTHS + 1)]


def prepare(df: pd.DataFrame, last_fact_date=None):
    """Дневная сетка до конца последнего прогнозного месяца + календарь.
    plan_actual / debt_actual — только известные значения (неизвестные дни пустые)."""
    cols = {DATE_COL: "date", FACT_COL: "fact", PLAN_COL: "plan"}
    if DEBT_COL:
        cols[DEBT_COL] = "debt"
    d = df[list(cols)].rename(columns=cols)
    d["date"] = pd.to_datetime(d["date"]).dt.normalize()
    value_cols = [c for c in d.columns if c != "date"]
    for c in value_cols:  # суммы из базы бывают Decimal/строками
        d[c] = pd.to_numeric(d[c], errors="coerce")
    d = d.groupby("date", as_index=False)[value_cols].sum(min_count=1)

    if last_fact_date is None:
        last_fact_date = d.loc[d["fact"].notna(), "date"].max()
    last_fact_date = pd.Timestamp(last_fact_date)
    cur = (last_fact_date + pd.Timedelta(days=1)).to_period("M")      # первый месяц без полного факта
    targets = target_months(cur)
    start = d["date"].min()
    # сетка — на месяц дальше последнего прогнозного: у последних дней есть соседние дни плана и «к оплате»
    end = max(d["date"].max(), (max(targets) + 1).end_time.normalize(), (cur + 2).end_time.normalize())

    d = pd.DataFrame({"date": pd.date_range(start, end, freq="D")}).merge(d, on="date", how="left")
    d["ym"] = d["date"].dt.to_period("M")
    d["mi"] = d["date"].dt.year * 12 + d["date"].dt.month      # номер месяца — для весов давности
    if CUMULATIVE_INPUT:  # накопленные с начала месяца -> дневные
        for c in ["fact", "plan"]:
            d[c] = d.groupby("ym")[c].transform(_cum_to_daily)
    kf = d["date"] <= last_fact_date
    d["fact"] = d["fact"].where(kf).fillna(0.0).where(kf).round(2)

    for c in value_cols[1:]:  # plan, debt
        last = d.loc[d[c].notna(), "date"].max()
        kn = d["date"] <= last
        # пропуск внутри известного: план и «к оплате» — 0 (на этот день ничего нет), остаток портфеля — как вчера
        filled = d[c].ffill() if (c == "debt" and RR_DENOM == "last") else d[c].fillna(0.0)
        d[f"{c}_actual"] = filled.where(kn).round(2)
    # месяц с нулевым планом считаем незагруженным
    zero = d.groupby("ym")["plan_actual"].transform("sum") == 0
    d.loc[zero, "plan_actual"] = np.nan

    d = d.merge(build_calendar(start, end), on="date", how="left")
    return d, last_fact_date, cur, targets


# ============================================================================
# ПРОГНОЗ ПЛАНА И ДОЛГА (для дней и месяцев, где их ещё нет)
# ============================================================================
def full_months(frame: pd.DataFrame, col: str) -> pd.Series:
    """Итоги месяцев, где ряд известен на все дни месяца."""
    g = frame.groupby("ym")[f"{col}_actual"]
    n_known, n_rows = g.count(), g.size()
    dim = pd.Series([m.days_in_month for m in n_rows.index], index=n_rows.index)
    ok = (n_known == n_rows) & (n_rows == dim)
    return g.sum()[ok]


def trend_forecast(s: pd.Series, T, h: float) -> float:
    """Последнее значение × сглаженный темп роста (h = 0 — без роста; иначе вес темпа падает вдвое каждые h мес.)."""
    s = s.dropna()
    s = s[s > 0]
    if s.empty:
        return np.nan
    steps = (T - s.index[-1]).n
    g = 0.0
    if h != 0 and len(s) >= 2:
        gaps = np.array([(b - a).n for a, b in zip(s.index[:-1], s.index[1:])], float)
        lg = np.diff(np.log(s.to_numpy())) / gaps
        ages = np.arange(len(lg))[::-1]
        w = 0.5 ** (ages / h) if np.isfinite(h) else np.ones(len(lg))
        g = (w * lg).sum() / w.sum()
    return float(s.iloc[-1] * np.exp(g * steps))


def series_total_forecast(totals: pd.Series, T, cfg) -> float:
    """Итог месяца T по полным месяцам до T: последний × темп роста (за месяц или в пересчёте на день)."""
    hist = totals[totals.index < T]
    if hist.empty:
        return np.nan
    if cfg["unit"] == "день":
        per_day = hist / np.array([m.days_in_month for m in hist.index])
        return trend_forecast(per_day, T, cfg["h"]) * T.days_in_month
    return trend_forecast(hist, T, cfg["h"])


def series_profile(frame: pd.DataFrame, col: str, T, how: str, totals: pd.Series) -> np.ndarray:
    """Раскладка итога месяца T по дням (доли, сумма = 1) по последним PROFILE_MONTHS полным месяцам до T:
    по числу месяца или по номеру рабочего дня (все нерабочие дни — одна общая доля)."""
    rows = frame[frame["ym"] == T]
    hist = [m for m in totals.index if m < T][-PROFILE_MONTHS:]
    if not hist:
        return np.full(len(rows), 1.0 / len(rows))
    h = frame[frame["ym"].isin(hist)]
    share = (h[f"{col}_actual"] / h["ym"].map(totals).replace(0, np.nan)).to_numpy()
    if how == "по рабочим дням":
        key_h = np.where(h["is_workday"] == 1, h["wd_num"], 0)
        key_t = np.where(rows["is_workday"] == 1, rows["wd_num"], 0)
    else:
        key_h, key_t = h["day"].to_numpy(), rows["day"].to_numpy()
    mean_by_key = pd.Series(share).groupby(key_h).mean()
    prof = pd.Series(key_t).map(mean_by_key).to_numpy(dtype=float)
    fallback = np.nanmean(share) if np.isfinite(share).any() else 1.0
    prof = np.where(np.isnan(prof), fallback, prof)
    s = prof.sum()
    return prof / s if s > 0 else np.full(len(rows), 1.0 / len(rows))


def fill_series(frame: pd.DataFrame, col: str, cfg: dict, months=None) -> pd.Series:
    """Заполняет неизвестные дни ряда col (plan / debt). Известная часть месяца остаётся как есть;
    остальное = итог месяца × раскладка по дням. Итог месяца — по известной части (если она ≥ 25% месяца
    по раскладке), иначе по тренду прошлых месяцев."""
    s = frame[f"{col}_actual"].copy()
    totals = full_months(frame, col)
    todo = sorted(m for m in frame.loc[s.isna(), "ym"].unique() if months is None or m in months)
    for T in todo:
        idx = frame.index[frame["ym"] == T]
        known = s.loc[idx].notna().to_numpy()
        prof = series_profile(frame, col, T, cfg["profile"], totals)
        k_share = prof[known].sum()
        if known.any() and k_share >= 0.25:
            total = s.loc[idx][known].sum() / k_share
        else:
            total = series_total_forecast(totals, T, cfg)
        s.loc[idx[~known]] = total * prof[~known]
    return s


def tune_series(frame: pd.DataFrame, col: str) -> dict:
    """Подбор прогноза итога месяца (за месяц / в день, темп роста) и раскладки по дням — по прошлым месяцам."""
    totals = full_months(frame, col)
    test = [m for m in totals.index if m >= pd.Period(BACKTEST_FROM, "M") and (totals.index < m).any()]
    if not test:
        test = list(totals.index[1:])
    best = None
    for unit in ("месяц", "день"):
        for h in TREND_HALFLIVES:
            cfg = dict(unit=unit, h=h)
            err = np.array([series_total_forecast(totals, m, cfg) / totals[m] - 1 for m in test])
            score = np.nanmean(np.abs(err))
            if best is None or score < best["total_err"] - 1e-12:
                best = dict(cfg, total_err=score)
    prof_err = {}
    for how in ("по дням месяца", "по рабочим дням"):
        errs = []
        for m in test:
            a = frame.loc[frame["ym"] == m, f"{col}_actual"].to_numpy()
            p = series_profile(frame, col, m, how, totals)
            errs.append(np.abs(p * a.sum() - a).sum() / a.sum())
        prof_err[how] = np.mean(errs)
    best["profile"] = min(prof_err, key=prof_err.get)
    best["profile_err"] = prof_err[best["profile"]]
    return best


# ============================================================================
# ПРИЗНАКИ
# ============================================================================
PLAN_FEATURES = ["plan_share", "plan_cum_share", "plan_next3", "plan_prev3", "plan_prev_nw"]
DEBT_FEATURES = ["debt_share", "debt_cum_share", "debt_prev1", "debt_next1", "debt_prev3", "debt_next3",
                 "debt_prev_nw", "debt_next_nw"]
CAL_FEATURES = (["dow", "day", "month", "dim", "days_to_eom", "is_workday", "is_holiday",
                 "wd_num", "wd_total", "wd_left", "wd_frac", "nonwork_before", "nonwork_after",
                 "first_wd_month", "last_wd_month", "ny_break", "days_since_pay"]
                + [f"pay_{p}" for p in PAYDAYS])
PLAN_LEVEL = ["rate_prev1", "rate_prev3", "rate_ly", "plan_growth"]
DEBT_LEVEL = ["rr_prev1", "rr_prev3", "rr_ly", "debt_growth"]
LEVEL_COLS = PLAN_LEVEL + DEBT_LEVEL + ["debt_cover"]


def dist_columns(col: str) -> list:
    if col == "plan":
        return [col, f"{col}_m"] + PLAN_FEATURES
    return [col, f"{col}_m"] + DEBT_FEATURES + [f"{col}_t{l:+d}" for l in DUE_LAGS]


def add_dist_features(f: pd.DataFrame, col: str) -> pd.DataFrame:
    """Раскладка ряда col (план или «к оплате») по дням: доля дня в сумме месяца, суммы на соседних днях и на нерабочих
    днях прямо до и после дня (их часто платят в ближайший рабочий). Ряд — известный или спрогнозированный."""
    v = f[col]
    f[f"{col}_m"] = f.groupby("ym")[col].transform("sum")
    pm = f[f"{col}_m"].replace(0, np.nan)
    f[f"{col}_share"] = v / pm
    f[f"{col}_cum_share"] = f.groupby("ym")[col].cumsum() / pm
    f[f"{col}_next3"] = sum(v.shift(-k) for k in (1, 2, 3)) / pm
    f[f"{col}_prev3"] = sum(v.shift(k) for k in (1, 2, 3)) / pm
    x, work = np.nan_to_num(v.to_numpy(dtype=float)), f["is_workday"].to_numpy()
    prev_nw, next_nw = np.zeros(len(f)), np.zeros(len(f))
    acc = 0.0
    for i in range(len(f)):              # сумма на нерабочих днях прямо перед днём
        prev_nw[i] = acc
        acc = 0.0 if work[i] else acc + x[i]
    acc = 0.0
    for i in range(len(f) - 1, -1, -1):  # и сразу после него
        next_nw[i] = acc
        acc = 0.0 if work[i] else acc + x[i]
    f[f"{col}_prev_nw"] = prev_nw / pm
    if col != "plan":
        f[f"{col}_prev1"] = v.shift(1) / pm
        f[f"{col}_next1"] = v.shift(-1) / pm
        f[f"{col}_next_nw"] = next_nw / pm
        for lag in DUE_LAGS:                 # «к оплате» на дату t + lag — для ожидания оплат по датам платежа
            f[f"{col}_t{lag:+d}"] = v.shift(-lag)
    return f


def month_table(d: pd.DataFrame, last_fact_date) -> pd.DataFrame:
    g = d.groupby("ym")
    mt = pd.DataFrame(index=pd.period_range(d["ym"].min(), d["ym"].max(), freq="M"))
    mt["plan_m"] = g["plan"].sum()
    mt["plan_known"] = g["plan_actual"].count() == g.size()
    mt["fact_m"] = g["fact"].sum()
    whole = mt.index.start_time >= d["date"].min()
    mt["complete"] = whole & (mt.index.end_time.normalize() <= last_fact_date)
    ok = mt["complete"] & mt["plan_known"]
    mt["rate"] = np.where(ok, mt["fact_m"] / mt["plan_m"].replace(0, np.nan), np.nan)
    if "debt" in d:
        mt["debt_m"] = g["debt"].sum()
        mt["debt_known"] = g["debt_actual"].count() == g.size()
        okd = mt["complete"] & mt["debt_known"]
        mt["rr"] = np.where(okd, mt["fact_m"] / mt["debt_m"].replace(0, np.nan), np.nan)
    return mt


def month_level_features(mt: pd.DataFrame, rate: pd.Series, rr: pd.Series = None) -> pd.DataFrame:
    """Признаки уровня месяца, известные на его начало: выполнение плана и recovery прошлых месяцев, рост плана
    и «к оплате», «к оплате» этого месяца к его плану."""
    f = pd.DataFrame(index=mt.index)
    f["rate_prev1"] = rate.shift(1)
    f["rate_prev3"] = rate.shift(1).rolling(3, min_periods=1).mean()
    f["rate_ly"] = rate.shift(12)
    f["plan_growth"] = mt["plan_m"] / mt["plan_m"].shift(1).rolling(3, min_periods=1).mean() - 1
    if "debt_m" in mt:
        rr = mt["rr"] if rr is None else rr
        f["rr_prev1"] = rr.shift(1)
        f["rr_prev3"] = rr.shift(1).rolling(3, min_periods=1).mean()
        f["rr_ly"] = rr.shift(12)
        f["debt_growth"] = mt["debt_m"] / mt["debt_m"].shift(1).rolling(3, min_periods=1).mean() - 1
        f["debt_cover"] = mt["debt_m"] / mt["plan_m"].replace(0, np.nan)
    return f


def attach_level(frame: pd.DataFrame, lvl: pd.DataFrame) -> pd.DataFrame:
    frame = frame.drop(columns=[c for c in LEVEL_COLS if c in frame.columns])
    return frame.merge(lvl, left_on="ym", right_index=True, how="left")


def daily_debt(d_or_mt) -> bool:
    """Модели от «к оплате» по дням — только если оно известно на весь месяц вперёд. Иначе в прогнозе раскладка
    «к оплате» по дням была бы прогнозом, а модель училась на фактической, — используем только месячные признаки."""
    return ("debt" in d_or_mt or "debt_m" in d_or_mt) and DEBT_KNOWN_AHEAD


def feature_sets(has_debt: bool, by_day: bool) -> dict:
    """Модели «к плану» (boost, cal) учат долю плана месяца на день; модели от «к оплате» (boost_d, cal_d) — поправку
    к ожиданию оплат по датам платежа (DueModel). boost — с признаками уровня, cal — без них (уровень — отдельно)."""
    debt_days = DEBT_FEATURES if by_day else []
    debt_month = DEBT_LEVEL + ["debt_cover"] if has_debt else []
    fs = {"boost": CAL_FEATURES + PLAN_FEATURES + debt_days + PLAN_LEVEL + debt_month,
          "cal": CAL_FEATURES + PLAN_FEATURES + debt_days}
    if by_day:
        fs["boost_d"] = CAL_FEATURES + DEBT_FEATURES + DEBT_LEVEL
        fs["cal_d"] = CAL_FEATURES + DEBT_FEATURES
    return fs


# ============================================================================
# МОДЕЛИ
# ============================================================================
def make_model(seed=RANDOM_STATE):
    try:
        import lightgbm as lgb
        return lgb.LGBMRegressor(
            n_estimators=700, learning_rate=0.02, num_leaves=15, min_child_samples=15,
            subsample=0.8, subsample_freq=1, colsample_bytree=0.8, reg_lambda=1.0,
            random_state=seed, verbose=-1)
    except ImportError:
        from sklearn.ensemble import HistGradientBoostingRegressor
        kw = dict(max_iter=700, learning_rate=0.03, max_leaf_nodes=15, min_samples_leaf=15,
                  l2_regularization=1.0, early_stopping=False, random_state=seed)
        try:
            return HistGradientBoostingRegressor(max_features=0.8, **kw)   # scikit-learn ≥ 1.4
        except TypeError:
            return HistGradientBoostingRegressor(**kw)


class SeedBag:
    """Среднее нескольких бустингов с разными seed: на коротком ряду это заметно снижает разброс прогноза."""

    def __init__(self, n=None):
        self.models = [make_model(RANDOM_STATE + i) for i in range(n or N_SEEDS)]

    def fit(self, X, y, sample_weight=None):
        for m in self.models:
            m.fit(X, y, sample_weight=sample_weight)
        if hasattr(self.models[0], "feature_importances_"):
            self.feature_importances_ = np.mean([m.feature_importances_ for m in self.models], axis=0)
        return self

    def predict(self, X):
        return np.mean([m.predict(X) for m in self.models], axis=0)


def _train_rows(train: pd.DataFrame, target_month, ycol: str):
    """Строки обучения (без EXCLUDE_DATES) и веса давности."""
    tr = train[~train["date"].isin(pd.to_datetime(EXCLUDE_DATES)) & train[ycol].notna()]
    w = None
    if RECENCY_HALFLIFE:
        age = (target_month.year * 12 + target_month.month - tr["mi"]).to_numpy(dtype=float)
        w = 0.5 ** (age / RECENCY_HALFLIFE)
    return tr, w


def fit_model(train: pd.DataFrame, target_month, feats: list, ycol: str = "y"):
    tr, w = _train_rows(train, target_month, ycol)
    feats = [f for f in feats if tr[f].notna().any()]  # признак без единого значения (напр. rate_ly в первый год) — пропускаем
    model = SeedBag().fit(tr[feats], tr[ycol], sample_weight=w)
    model.feats_ = feats
    return model


class DueKernel:
    """Ожидаемые оплаты дня по датам платежа: E(t) = Σ_l w_l · «к оплате»(t + l) + w_0 · «к оплате» месяца / дней.
    w_l ≥ 0 — какую долю платят за l дней до даты платежа (l > 0) или через −l дней после (l < 0); оцениваются по
    прошлым месяцам. Ожидание линейно по «к оплате»: всплеск больше виденного в истории (первый рабочий день после
    длинных праздников) масштабируется, а не обрезается, как в деревьях."""

    @staticmethod
    def _X(rows: pd.DataFrame) -> np.ndarray:
        cols = [rows[f"debt_t{lag:+d}"].to_numpy(dtype=float) for lag in DUE_LAGS]
        cols.append((rows["debt_m"] / rows["dim"]).to_numpy(dtype=float))
        return np.nan_to_num(np.column_stack(cols))

    def fit(self, rows: pd.DataFrame, sample_weight=None):
        from sklearn.linear_model import LinearRegression
        self.lr = LinearRegression(positive=True, fit_intercept=False).fit(
            self._X(rows), rows["fact"].to_numpy(dtype=float), sample_weight=sample_weight)
        return self

    def expected(self, rows: pd.DataFrame) -> np.ndarray:
        floor = 0.02 * np.nan_to_num((rows["debt_m"] / rows["dim"]).to_numpy(dtype=float))
        return np.maximum(self._X(rows) @ self.lr.coef_, 0) + np.maximum(floor, 1e-6)

    def weights(self) -> pd.Series:
        """Веса по сдвигу от даты платежа, доли от суммы (без равномерной части)."""
        w = pd.Series(self.lr.coef_[:-1], index=list(DUE_LAGS))
        return w / w.sum() if w.sum() > 0 else w


class DueModel:
    """Модель от «к оплате»: ожидание оплат по датам платежа E(t) × поправка бустинга r(t) = факт / E(t)
    (календарь, выходные и праздники, дни зарплаты, recovery прошлых месяцев)."""

    def __init__(self, feats: list):
        self.feats = feats

    @staticmethod
    def _with_share(rows: pd.DataFrame, E: np.ndarray) -> pd.DataFrame:
        return rows.assign(debt_expected_share=E / rows["debt_m"].replace(0, np.nan).to_numpy())

    def fit(self, rows: pd.DataFrame, sample_weight=None, kernel: DueKernel = None):
        self.kernel = kernel or DueKernel().fit(rows, sample_weight)
        E = self.kernel.expected(rows)
        X = self._with_share(rows, E)
        self.feats_ = [f for f in self.feats + ["debt_expected_share"] if X[f].notna().any()]
        self.bag = SeedBag().fit(X[self.feats_], rows["fact"].to_numpy(dtype=float) / E, sample_weight=sample_weight)
        self.feature_importances_ = getattr(self.bag, "feature_importances_", None)
        return self

    def predict_money(self, rows: pd.DataFrame) -> np.ndarray:
        E = self.kernel.expected(rows)
        return np.clip(self.bag.predict(self._with_share(rows, E)[self.feats_]), 0, None) * E


def fit_models(dd: pd.DataFrame, mt: pd.DataFrame, until, fs: dict) -> dict:
    """Все модели, обученные на полных месяцах до until: «к плану» и, если есть «к оплате», от него."""
    ok_plan = mt.index[mt["complete"] & mt["plan_known"]]
    tr = dd[(dd["ym"] < until) & dd["ym"].isin(ok_plan)]
    models = {k: fit_model(tr, until, fs[k], "y") for k in ("boost", "cal")}
    if "boost_d" in fs:
        ok_debt = mt.index[mt["complete"] & mt["debt_known"]]
        trd, w = _train_rows(dd[(dd["ym"] < until) & dd["ym"].isin(ok_debt)], until, "y_debt")
        kernel = DueKernel().fit(trd, w)
        for k in ("boost_d", "cal_d"):
            models[k] = DueModel(fs[k]).fit(trd, w, kernel)
    return models


def predict_share(model, rows: pd.DataFrame) -> np.ndarray:
    """Прогноз доли месячного плана (или «к оплате») на каждый день."""
    return np.clip(model.predict(rows[model.feats_]), 0, None)


# ============================================================================
# РЕФЕРЕНСНЫЕ МЕСЯЦЫ (текущий подход)
# ============================================================================
def reference_curve(d: pd.DataFrame, month, refs) -> np.ndarray:
    """Текущий подход: средняя кривая выполнения (накопл. факт / план месяца) по референсным месяцам,
    по календарному дню. Возвращает дневные доли плана (умножить на план месяца)."""
    n = month.days_in_month
    curves = []
    for r in refs:
        g = d[d["ym"] == r]
        if g.empty or g["fact"].isna().any() or not g["plan_m"].iloc[0] > 0:
            continue
        c = (g["fact"].cumsum() / g["plan_m"]).to_numpy()
        curves.append(np.concatenate([c, np.repeat(c[-1], max(0, n - len(c)))])[:n])
    if not curves:
        return np.full(n, np.nan)
    return np.diff(np.r_[0.0, np.mean(curves, axis=0)])


def _wd_keys(g: pd.DataFrame) -> np.ndarray:
    """Рабочий день — его номер от начала месяца (первая половина рабочих дней) или от конца (-1 — последний);
    все нерабочие дни — 0."""
    wd, left, tot = g["wd_num"].to_numpy(), g["wd_left"].to_numpy(), g["wd_total"].to_numpy()
    k = np.where(wd <= (tot + 1) // 2, wd, -(left + 1))
    return np.where(g["is_workday"].to_numpy() == 1, k, 0)


def reference_wd(d: pd.DataFrame, month, refs) -> np.ndarray:
    """Референсные месяцы с выравниванием по рабочим дням: доля плана рабочего дня — как у рабочего дня с тем же
    номером от начала или от конца месяца, нерабочего — средняя по нерабочим дням. Итог месяца — средняя
    эффективность референсных месяцев. Возвращает дневные доли плана."""
    rows = d[d["ym"] == month]
    shares, totals = [], []
    for r in refs:
        g = d[d["ym"] == r]
        if g.empty or g["fact"].isna().any() or not g["plan_m"].iloc[0] > 0:
            continue
        s = g["fact"].to_numpy() / g["plan_m"].iloc[0]
        shares.append(pd.Series(s).groupby(_wd_keys(g)).mean())
        totals.append(s.sum())
    if not shares:
        return np.full(len(rows), np.nan)
    by_key = pd.concat(shares, axis=1).mean(axis=1)
    prof = pd.Series(_wd_keys(rows)).map(by_key).to_numpy(dtype=float)
    work = by_key[by_key.index != 0]
    prof = np.where(np.isnan(prof), work.mean() if len(work) else np.nanmean(prof), prof)
    return prof / prof.sum() * np.mean(totals)


# ============================================================================
# СЦЕНАРИИ: КАКИМ МЕСЯЦ ВИДЕН НА ЕГО НАЧАЛО
# ============================================================================
SCENARIOS = {"": "план загружен", "_pu": "план ещё не загружен"}


def scenario_rows(d, mt, m, hide_plan: bool, hide_debt: bool, plan_cfg: dict, debt_cfg: dict) -> pd.DataFrame:
    """Строки месяца m такими, какими они видны на его начало: план и/или «к оплате» месяца m (и дальше)
    ещё не известны и спрогнозированы только по месяцам до m."""
    rows = d[d["ym"] == m].copy()
    mt2 = mt.copy()
    for col, hide, cfg in (("plan", hide_plan, plan_cfg), ("debt", hide_debt, debt_cfg)):
        if not hide or col not in d:
            continue
        f = d[["date", "ym", "day", "is_workday", "wd_num", f"{col}_actual"]].copy()
        f.loc[f["ym"] >= m, f"{col}_actual"] = np.nan
        f[col] = fill_series(f, col, cfg, months={m, m + 1})
        f = add_dist_features(f, col)
        fm = f[f["ym"] == m]
        for c in dist_columns(col):
            rows[c] = fm[c].to_numpy()
        mt2.loc[m, f"{col}_m"] = rows[f"{col}_m"].iloc[0]
    return attach_level(rows, month_level_features(mt2, mt2["rate"]))


def month_inputs(rows: pd.DataFrame, models: dict, d: pd.DataFrame, T, refs, sfx: str, M: dict = None) -> dict:
    """Составляющие прогноза месяца T в сценарии sfx: прогнозы и доли моделей, референсные кривые."""
    M = {} if M is None else M
    P = rows["plan_m"].iloc[0]
    M.update({"ym": T, "n": len(rows), "plan" + sfx: rows["plan"].to_numpy(), "P" + sfx: P})
    M["boost" + sfx] = predict_share(models["boost"], rows) * P
    M["cal" + sfx] = predict_share(models["cal"], rows)
    if "boost_d" in models:
        D = rows["debt_m"].iloc[0]
        M["D" + sfx] = D
        M["boost_d" + sfx] = models["boost_d"].predict_money(rows)
        M["cal_d" + sfx] = models["cal_d"].predict_money(rows) / D           # доля «к оплате» месяца
    if "ref" not in M:
        M["ref"] = reference_curve(d, T, refs)
        M["ref_wd"] = reference_wd(d, T, refs)
    return M


# ============================================================================
# NOWCAST: ПОПРАВКА ОСТАТКА ПО ФАКТУ МЕСЯЦА
# ============================================================================
NC_ARR = np.array([(0.0, 0.0)] + [(a, t) for a in NOWCAST_MAX for t in NOWCAST_TAU])  # (α_max, τ)


def nowcast_alpha(pred: np.ndarray, nc) -> np.ndarray:
    """Сила поправки остатка [конфигурация, k] при k = 0..n-1 известных днях: α_max · c / (c + τ),
    c — доля месяца, которая по прогнозу должна была прийти за первые k дней."""
    nc = np.atleast_2d(nc)
    cp = np.r_[0.0, np.cumsum(pred)][:-1]
    tot = pred.sum()
    c = cp / tot if tot > 0 else np.zeros_like(cp)
    with np.errstate(divide="ignore", invalid="ignore"):
        w = np.where(c > 0, c / (c + nc[:, 1:]), 0.0)
    return nc[:, :1] * w


def eom_forecasts(actual: np.ndarray, pred: np.ndarray, alphas: np.ndarray) -> np.ndarray:
    """Прогноз итога месяца [конфигурация, k] при известных первых k днях (векторно).
    Остаток прогноза домножается на 1 + α·(факт / прогноз за известные дни − 1)."""
    ca = np.r_[0.0, np.cumsum(actual)][:-1]   # факт за первые k дней, k = 0..n-1
    cp = np.r_[0.0, np.cumsum(pred)][:-1]
    rem = pred.sum() - cp
    ratio = np.where(cp > 0, ca / np.where(cp > 0, cp, 1), 1.0)
    ratio[0] = 1.0
    return ca + rem * (1 + alphas * (ratio - 1))


def eff_errors(M: dict, pred: np.ndarray, sfx: str, nc=NC_ARR) -> np.ndarray:
    """Разница эффективности прогноза и факта на конец месяца, п.п. [конфигурация поправки, k]:
    прогноз / план (загруженный или спрогнозированный) − факт / план."""
    a = M["fact"]
    return 100 * (eom_forecasts(a, pred, nowcast_alpha(pred, nc)) / M["P" + sfx] - a.sum() / M["P"])


def nowcast_paths(actual: np.ndarray, pred: np.ndarray, nc) -> np.ndarray:
    """Накопленный факт + прогноз остатка с поправкой [k, день] при k = 0..n-1 известных днях."""
    n = len(actual)
    alpha = nowcast_alpha(pred, nc)[0]
    ca = np.cumsum(actual)
    cums = np.empty((n, n))
    for k in range(n):
        ka, kp = (ca[k - 1] if k else 0.0), pred[:k].sum()
        ratio = ka / kp if k and kp > 0 else 1.0
        cums[k] = np.r_[ca[:k], ka + np.cumsum(pred[k:] * (1 + alpha[k] * (ratio - 1)))]
    return cums


def path_rows(actual: np.ndarray, cums: np.ndarray, plan_true: np.ndarray, plan_used: np.ndarray) -> list:
    """Ошибки по путям nowcast_paths при k = 0..n-1: разница эффективности на конец месяца и на каждую дату
    после k (накопленный прогноз / накопленный план против накопленного факта / накопленного плана), п.п."""
    n = len(actual)
    ca, cpt, cpu = np.cumsum(actual), np.cumsum(plan_true), np.cumsum(plan_used)
    p0 = np.diff(np.r_[0.0, cums[0]])           # прогноз без факта месяца
    day = np.arange(1, n + 1)
    with np.errstate(divide="ignore", invalid="ignore"):
        eff_fact = ca / cpt
        diffs = 100 * (cums / cpu - eff_fact)    # [k, день]
    rows = []
    for k in range(n):
        dk = diffs[k, k:]
        ok = np.isfinite(dk)
        ok5 = ok & (day[k:] >= 5)
        rows.append(dict(
            known_days=k,
            eff_fact=100 * ca[-1] / cpt[-1],                                     # эффективность факта на конец, %
            err_eff_pp=100 * (cums[k, -1] / cpu[-1] - ca[-1] / cpt[-1]),         # разница на конец месяца, п.п.
            dates_pp=np.abs(dk[ok]).mean() if ok.any() else np.nan,              # ср. |разница| на датах после k
            dates_pp_5=np.abs(dk[ok5]).mean() if ok5.any() else np.nan,          # то же с 5-го числа
            dates_within1=(np.abs(dk[ok]) <= 1).mean() if ok.any() else np.nan,
            err_eom=cums[k, -1] / ca[-1] - 1,                                    # относительная ошибка суммы
            wape=np.abs(p0[k:] - actual[k:]).sum() / max(actual[k:].sum(), 1e-9)))
    return rows


# ============================================================================
# УРОВЕНЬ ДЛЯ «КАЛЕНДАРЬ × УРОВЕНЬ / RECOVERY»
# ============================================================================
def level_estimate(L: pd.Series, T, h: float, w0: float) -> float:
    """Уровень на месяц T = взвешенное среднее уровней прошлых месяцев (вес падает вдвое каждые h мес.)
    и «нейтрального» уровня 1 с весом w0. Уровень месяца = факт / ожидание модели без признаков уровня,
    то есть выполнение (или recovery), очищенное от календаря."""
    if np.isinf(w0):
        return 1.0
    hist = L[L.index < T].dropna()
    ages = np.array([(T - m).n for m in hist.index], float)
    w = 0.5 ** (ages / h) if np.isfinite(h) else np.ones(len(ages))
    den = w0 + w.sum()
    return float((w0 + (w * hist.to_numpy()).sum()) / den) if den > 0 else 1.0


def histories(months: list) -> dict:
    """Уровни месяцев бэктеста (каждый — по моделям, обученным только до этого месяца): L_plan — факт / ожидание
    календарной модели к плану, L_debt — то же от «к оплате», A — факт / календарный фактор (для тренда факта)."""
    idx = pd.PeriodIndex([M["ym"] for M in months], freq="M")
    F = np.array([M["fact"].sum() for M in months])
    S = np.array([M["cal"].sum() for M in months])
    P = np.array([M["P"] for M in months])
    h = {"L_plan": pd.Series(F / (S * P), index=idx), "A": pd.Series(F / S, index=idx)}
    if "cal_d" in months[0]:
        E = np.array([M["cal_d"].sum() * M["D"] for M in months])
        h["L_debt"] = pd.Series(F / E, index=idx)
    return h


# ============================================================================
# БЭКТЕСТ
# ============================================================================
def backtest(d: pd.DataFrame, mt: pd.DataFrame, plan_cfg: dict, debt_cfg: dict = None) -> list:
    """Для каждого месяца с BACKTEST_FROM все модели обучены только на месяцах до него, а план и «к оплате» —
    такие, какими были известны на 1-е число (план — загружен или ещё нет)."""
    has_debt = "debt" in d
    fs = feature_sets(has_debt, daily_debt(d))
    dd = attach_level(d, month_level_features(mt, mt["rate"]))
    ok_plan = mt.index[mt["complete"] & mt["plan_known"]]
    ok_debt = mt.index[mt["complete"] & mt["debt_known"]] if has_debt else ok_plan
    start = pd.Period(BACKTEST_FROM, "M")
    test = [m for m in ok_plan if m >= start and m in ok_debt and (ok_debt < m).sum() >= 6]
    if len(test) < 3:
        raise ValueError("Мало полных месяцев для бэктеста — сдвиньте BACKTEST_FROM раньше")
    hide_debt = has_debt and not DEBT_KNOWN_AHEAD
    months = []
    for m in test:
        models = fit_models(dd, mt, m, fs)
        refs = [m - i for i in range(1, N_REF_MONTHS + 1)]
        M = None
        for s in SCENARIOS:
            rows = scenario_rows(d, mt, m, s == "_pu", hide_debt, plan_cfg, debt_cfg)
            M = month_inputs(rows, models, dd, m, refs, s, M)
        g = d[d["ym"] == m]
        M.update(date=g["date"].to_numpy(), fact=g["fact"].to_numpy())
        if has_debt:
            M["D_true"] = g["debt"].sum()
        months.append(M)
        print(f"  бэктест {m}: обучено на {dd.loc[(dd['ym'] < m) & dd['ym'].isin(ok_plan), 'ym'].nunique()} мес.")
    return months


# ============================================================================
# МЕТОДЫ, ВЫБОР И ЧЕСТНАЯ ОЦЕНКА
# ============================================================================
NAMES = {"ref": "Референсные месяцы", "ref_wd": "Реф. месяцы по раб. дням", "boost": "Бустинг к плану",
         "cal": "Календарь × уровень к плану", "trend": "Календарь × тренд факта",
         "boost_d": "Бустинг от «к оплате»", "cal_d": "Календарь × recovery", "mix": "Смесь", "auto": "Авто"}
LEVEL_CFGS = list(itertools.product(LEVEL_HALFLIVES, LEVEL_PRIOR_WEIGHTS))


def method_list(by_day_debt: bool) -> list:
    return ["ref", "ref_wd", "boost", "cal", "trend"] + (["boost_d", "cal_d"] if by_day_debt else [])


def method_cfgs(meth: str) -> list:
    if meth in ("cal", "cal_d"):
        return LEVEL_CFGS
    if meth == "trend":
        return list(TREND_HALFLIVES)
    return [None]


def method_pred(M: dict, meth: str, cfg, sfx: str, hist: dict) -> np.ndarray:
    """Прогноз метода на каждый день месяца без факта этого месяца, в деньгах."""
    T, P = M["ym"], M["P" + sfx]
    if meth in ("ref", "ref_wd"):
        return np.nan_to_num(M[meth]) * P
    if meth in ("boost", "boost_d"):
        return M[meth + sfx]
    if meth == "cal":
        return M["cal" + sfx] * P * level_estimate(hist["L_plan"], T, *cfg)
    if meth == "cal_d":
        return M["cal_d" + sfx] * M["D" + sfx] * level_estimate(hist["L_debt"], T, *cfg)
    if meth == "trend":   # факт / календарный фактор прошлых месяцев × тренд; без истории — уровень 1 к плану
        a = trend_forecast(hist["A"][hist["A"].index < T], T, cfg)
        return M["cal" + sfx] * (a if np.isfinite(a) else P)
    raise ValueError(meth)


def config_pred(ch: dict, M: dict, sfx: str, hist: dict) -> np.ndarray:
    """Прогноз выбранного варианта — метода или смеси."""
    if ch["method"] == "mix":
        return sum(w * method_pred(M, t, c, sfx, hist) for w, (t, c) in zip(ch["w"], ch["parts"]))
    return method_pred(M, ch["method"], ch["cfg"], sfx, hist)


def score_table(months: list, hist: dict, methods: list, sfx: str):
    """Прогнозы всех методов и конфигураций на все месяцы бэктеста и их средняя по k |разница эффективности|
    на конец месяца: SC[метод] — [месяц, конфигурация, поправка]."""
    SC, PR = {}, {}
    for meth in methods:
        PR[meth] = [[method_pred(M, meth, c, sfx, hist) for M in months] for c in method_cfgs(meth)]
        SC[meth] = np.array([[np.abs(eff_errors(M, PR[meth][ci][mi], sfx)).mean(axis=1)
                              for ci in range(len(PR[meth]))] for mi, M in enumerate(months)])
    return SC, PR


def weight_grid(k: int, step: float = 0.25) -> list:
    n = int(round(1 / step))
    return [tuple(x * step for x in c) for c in itertools.product(range(n + 1), repeat=k) if sum(c) == n and max(c) < n]


def select(SC: dict, PR: dict, months: list, idx: list, methods: list, sfx: str):
    """Выбор по месяцам idx: лучшая конфигурация каждого метода (с поправкой по факту), смесь лучших методов и
    победитель — по средней |разнице эффективности на конец месяца| по всем k."""
    best = {}
    for meth in methods:
        s = SC[meth][idx].mean(axis=0)
        ci, ni = np.unravel_index(np.nanargmin(s), s.shape)
        best[meth] = dict(method=meth, cfg=method_cfgs(meth)[ci], ci=ci, nc=NC_ARR[ni], score=float(s[ci, ni]))
    top = sorted(methods, key=lambda x: best[x]["score"])[:MIX_TOP]
    for w in weight_grid(len(top)):
        err = np.mean([np.abs(eff_errors(months[i], sum(wk * PR[t][best[t]["ci"]][i] for wk, t in zip(w, top)),
                                         sfx)).mean(axis=1) for i in idx], axis=0)
        j = int(np.nanargmin(err))
        if "mix" not in best or err[j] < best["mix"]["score"]:
            best["mix"] = dict(method="mix", w=w, parts=[(t, best[t]["cfg"]) for t in top], nc=NC_ARR[j],
                               score=float(err[j]))
    winner = min(best, key=lambda x: best[x]["score"])
    return best, winner


def honest(SC: dict, PR: dict, months: list, hist: dict, methods: list, sfx: str) -> dict:
    """Честный бэктест: месяц i прогнозируется вариантами, выбранными по месяцам 0..i-1 (первые HONEST_WARMUP
    месяцев — только для выбора). «Авто» — победитель выбора, как в рабочем прогнозе."""
    rows, paths, picks = [], {}, {}
    for i in range(min(max(HONEST_WARMUP, 1), len(months) - 1), len(months)):
        best, winner = select(SC, PR, months, list(range(i)), methods, sfx)
        M = months[i]
        for name in list(best) + ["auto"]:
            ch = best[winner if name == "auto" else name]
            cums = nowcast_paths(M["fact"], config_pred(ch, M, sfx, hist), ch["nc"])
            for r in path_rows(M["fact"], cums, M["plan"], M["plan" + sfx]):
                rows.append(dict(method=name, ym=str(M["ym"]), **r))
            if name == "auto":
                paths[M["ym"]], picks[M["ym"]] = cums, winner
    return dict(df=pd.DataFrame(rows), paths=paths, picks=picks)


# ============================================================================
# ПОЛНЫЙ ПРОГНОЗ НА МЕСЯЦЫ (по дням)
# ============================================================================
def add_effs(o: pd.DataFrame) -> pd.DataFrame:
    """Накопленные суммы и эффективность на дату по дневной таблице (план, факт, прогнозы)."""
    g = o.groupby("month", sort=False)
    o["cum_plan"] = g["plan"].cumsum()
    o["cum_fact"] = g["fact"].cumsum().where(o["has_fact"])
    o["cum_forecast"] = g["forecast"].cumsum()
    o["cum_fact_plus_forecast"] = g["fact_plus_forecast"].cumsum()
    cp = o["cum_plan"].replace(0, np.nan)
    o["eff_fact_pct"] = 100 * o["cum_fact"] / cp                                    # эффективность факта на дату
    o["eff_forecast_pct"] = 100 * o["cum_forecast"] / cp                            # эффективность прогноза на дату
    o["eff_fact_plus_forecast_pct"] = 100 * o["cum_fact_plus_forecast"] / cp        # куда придёт месяц
    o["eff_diff_pp"] = o["eff_forecast_pct"] - o["eff_fact_pct"]                    # проверка по факту, п.п.
    o["forecast_err_pct"] = 100 * (o["cum_forecast"] / o["cum_fact"] - 1)
    return o


def _interval(h: pd.DataFrame, k: int, n: int, eff: float):
    """80% интервал эффективности на конец месяца по честному бэктесту «Авто» при k известных днях."""
    if h is None or h.empty or k >= n:
        return np.nan, np.nan
    a = h[h["method"] == "auto"]
    e = a.loc[a["known_days"] == min(k, a["known_days"].max()), "err_eff_pp"]
    if len(e) < 5:
        return np.nan, np.nan
    return eff - e.quantile(0.9), eff - e.quantile(0.1)


def _fmt(x):
    return f"{x:,.0f}".replace(",", " ")


def month_summary(o: pd.DataFrame, h: pd.DataFrame, plan_forecast: bool, method: str) -> dict:
    """Строка сводки по месяцу: эффективность прогноза без факта, факт + прогноз остатка, интервал."""
    n, k = len(o), int(o["has_fact"].sum())
    P = o["plan"].sum()
    eff, eff_pure = 100 * o["fact_plus_forecast"].sum() / P, 100 * o["forecast"].sum() / P
    lo, hi = _interval(h, k, n, eff)
    return {"месяц": o["month"].iloc[0], "дней факта": f"{k} из {n}",
            "план": _fmt(P) + (" (прогноз)" if plan_forecast else ""),
            "эфф.: прогноз без факта": f"{eff_pure:.1f}%",
            "эфф.: факт + прогноз остатка": f"{eff:.1f}%",
            "80% интервал": f"{lo:.1f}% – {hi:.1f}%" if np.isfinite(lo) else "—",
            "сумма: факт + прогноз остатка": _fmt(o["fact_plus_forecast"].sum()),
            "разница прогноза с фактом на сегодня": f"{o['eff_diff_pp'].iloc[k - 1]:+.2f} п.п." if k else "—",
            "метод": method}


def production(d, mt, cur, targets, sel, hist, hon):
    """Прогноз на месяцы targets по каждому дню. Для текущего и следующих месяцев модели обучены на всех полных
    месяцах, для прошедших — только на месяцах до них (как прогноз, сделанный тогда)."""
    has_debt = "debt" in d
    fs = feature_sets(has_debt, daily_debt(d))
    ok_plan = mt.index[mt["complete"] & mt["plan_known"]]
    rate = mt["rate"].copy()
    rr = mt["rr"].copy() if has_debt else None
    hist = {k: v.copy() for k, v in hist.items()}
    series = d["fact"].copy()  # факт; по мере прогноза дописываем прогноз (для recovery с начала данных)
    months = sorted(set(t for t in targets if t < cur) |
                    (set(pd.period_range(cur, max(targets), freq="M")) if max(targets) >= cur else set()))
    models, out, summary, dev = {}, [], [], {}

    for T in months:
        train_until = T if T < cur else cur   # прошедший месяц — модели, обученные только на месяцах до него
        dd = attach_level(d, month_level_features(mt, rate, rr))
        if train_until not in models:
            models[train_until] = fit_models(dd, mt, train_until, fs)
        rows = dd[dd["ym"] == T].copy()
        sc = "" if mt.loc[T, "plan_known"] else "_pu"
        best, winner = sel[sc]
        ch = best[winner]
        refs = [m for m in ok_plan if m < min(T, cur)][-N_REF_MONTHS:]
        M = month_inputs(rows, models[train_until], d, T, refs, sc)
        p = config_pred(ch, M, sc, hist)     # прогноз без факта месяца
        a = rows["fact"].to_numpy()
        n, k = len(a), int(np.sum(~np.isnan(a)))
        P = M["P" + sc]
        alpha = nowcast_alpha(p, ch["nc"])[0]
        ka, kp = np.nansum(a[:k]), p[:k].sum()
        ratio = ka / kp if k and kp > 0 else 1.0
        cum = np.r_[np.cumsum(a[:k]), ka + np.cumsum(p[k:] * (1 + (alpha[k] if k < n else 0.0) * (ratio - 1)))]
        fc = np.diff(np.r_[0.0, cum])
        series.loc[(d["ym"] == T).to_numpy()] = fc
        if T >= cur:   # уровни следующих месяцев — с учётом прогноза этого
            S = M["cal" + sc].sum()
            rate.loc[T] = cum[-1] / P
            hist["L_plan"].loc[T] = cum[-1] / (S * P)
            hist["A"].loc[T] = cum[-1] / S
            if has_debt:
                rr.loc[T] = cum[-1] / rows["debt_m"].iloc[0]
            if "L_debt" in hist:
                hist["L_debt"].loc[T] = cum[-1] / (M["cal_d" + sc].sum() * M["D" + sc])

        o = pd.DataFrame({
            "month": str(T), "date": rows["date"].to_numpy(), "is_workday": rows["is_workday"].to_numpy(),
            "has_fact": np.arange(n) < k,                           # True — по этому дню факт уже есть
            "plan": rows["plan"].to_numpy(), "plan_is_forecast": rows["plan_actual"].isna().to_numpy(),
            "fact": a,
            "forecast": p,                  # прогноз на КАЖДЫЙ день месяца, без единого факта этого месяца
            "fact_plus_forecast": fc})      # факт, где он есть, + прогноз остатка с поправкой по факту на сегодня
        o = add_effs(o)
        s = month_summary(o, hon[sc]["df"], bool(sc), NAMES[winner])

        if has_debt:
            debt = rows["debt"].to_numpy()
            hist_mask = (d["date"] < rows["date"].iloc[0]).to_numpy()
            f_before, d_before = series[hist_mask].sum(), d.loc[hist_mask, "debt"].sum()
            cum_debt = np.cumsum(debt) if RR_DENOM == "sum" else debt
            tot_debt = d_before + np.cumsum(debt) if RR_DENOM == "sum" else debt
            o["debt"] = debt
            o["debt_is_forecast"] = rows["debt_actual"].isna().to_numpy()
            with np.errstate(divide="ignore", invalid="ignore"):
                pos = debt > 0
                o["rr_day_fact_pct"] = np.where(pos, 100 * a / debt, np.nan)          # за день: факт / «к оплате»
                o["rr_day_forecast_pct"] = np.where(pos, 100 * p / debt, np.nan)      # за день: прогноз / «к оплате»
                o["rr_fact_pct"] = 100 * o["cum_fact"].to_numpy() / cum_debt          # с начала месяца
                o["rr_forecast_pct"] = 100 * o["cum_forecast"].to_numpy() / cum_debt
                o["rr_fact_plus_forecast_pct"] = 100 * cum / cum_debt
                o["rr_total_pct"] = 100 * (f_before + cum) / tot_debt                 # с начала данных
            note = " («к оплате» — прогноз)" if o["debt_is_forecast"].any() else ""
            s["recovery: прогноз без факта"] = f"{100 * o['cum_forecast'].iloc[-1] / cum_debt[-1]:.2f}%"
            s["recovery: факт + прогноз остатка"] = f"{100 * cum[-1] / cum_debt[-1]:.2f}%" + note
        if T in targets:
            out.append(o)
            summary.append(s)
            # прогноз итога месяца, сделанный при k = 0, 1, … известных днях (по сегодня) — для Excel
            dev[T] = nowcast_paths(np.nan_to_num(a), p, ch["nc"])[:min(k, n - 1) + 1, -1]
    return pd.concat(out, ignore_index=True), pd.DataFrame(summary), models, dev


# ============================================================================
# ГРАФИКИ
# ============================================================================
INK, INK2, GRID, SURFACE = "#0b0b0b", "#52514e", "#e4e3df", "#fcfcfb"
BLUE, ORANGE, AQUA = "#2a78d6", "#eb6834", "#1baf7a"


def _style(ax):
    ax.set_facecolor(SURFACE)
    for s in ["top", "right"]:
        ax.spines[s].set_visible(False)
    for s in ["left", "bottom"]:
        ax.spines[s].set_color(GRID)
    ax.tick_params(colors=INK2, labelsize=9)
    ax.grid(axis="y", color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)


def plot_backtest(curves: dict, path: str):
    """curves: {подпись: (Series средней |разницы| по known_days, цвет, стиль линии)}"""
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(9, 4.5), facecolor=SURFACE)
    _style(ax)
    for name, (s, color, ls) in curves.items():
        ax.plot(s.index, s.values, color=color, linewidth=2, linestyle=ls, label=name)
    ax.axhline(1, color=INK2, linewidth=1, linestyle=(0, (1, 2)))
    ax.annotate("цель 1 п.п.", (0, 1), xytext=(2, 4), textcoords="offset points", color=INK2, fontsize=9)
    ax.set_xlabel("Сколько дней месяца уже известно", color=INK2)
    ax.set_ylabel("Средняя |разница эффективности|, п.п.", color=INK2)
    ax.set_title("Честный бэктест: эффективность прогноза против факта на конец месяца",
                 color=INK, loc="left", fontsize=12)
    ax.legend(frameon=False, fontsize=9, labelcolor=INK2)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def plot_month(g: pd.DataFrame, path: str):
    """Эффективность на дату (и recovery rate, если есть «к оплате»): факт; прогноз на весь месяц без факта
    с коридором ±1 п.п. (цель); куда придёт месяц с учётом факта по сегодня."""
    import matplotlib.dates as mdates
    import matplotlib.pyplot as plt
    panels = [("eff", "Эффективность на дату, %", True)]
    if "rr_fact_pct" in g:
        panels.append(("rr", "Recovery rate с начала месяца, %", False))
    cols = {"eff": ("eff_fact_pct", "eff_forecast_pct", "eff_fact_plus_forecast_pct"),
            "rr": ("rr_fact_pct", "rr_forecast_pct", "rr_fact_plus_forecast_pct")}
    fig, axes = plt.subplots(1, len(panels), figsize=(6.5 * len(panels), 4.5), facecolor=SURFACE)
    axes = np.atleast_1d(axes)
    known, rest = g[g["has_fact"]], g[~g["has_fact"]]
    for ax, (key, title, band) in zip(axes, panels):
        c_fact, c_fc, c_fpf = cols[key]
        _style(ax)
        if band:
            ax.fill_between(g["date"], g[c_fc] - 1, g[c_fc] + 1, color=BLUE, alpha=0.15, linewidth=0,
                            label="±1 п.п. от прогноза (цель)")
        ax.plot(g["date"], g[c_fc], color=BLUE, linewidth=2, linestyle=(0, (4, 3)), label="Прогноз (без факта месяца)")
        if len(known):
            ax.plot(known["date"], known[c_fact], color=INK, linewidth=2, label="Факт")
        if len(rest):
            r = pd.concat([known.tail(1), rest]) if len(known) else rest
            ax.plot(r["date"], r[c_fpf], color=ORANGE, linewidth=2, linestyle=(0, (1, 1.5)),
                    label="Факт + прогноз остатка")
        if key == "eff":
            ax.axhline(100, color=INK2, linewidth=1, linestyle=(0, (1, 2)))
        ax.set_xlim(g["date"].min() - pd.Timedelta(hours=12), g["date"].max() + pd.Timedelta(hours=12))
        ax.xaxis.set_major_locator(mdates.DayLocator(bymonthday=[1, 8, 15, 22, 29]))
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%d.%m"))
        ax.set_title(f"{g['month'].iloc[0]}: {title}", color=INK, loc="left", fontsize=11)
        ax.legend(frameon=False, fontsize=8, labelcolor=INK2, loc="upper right" if key == "eff" else "lower right")
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


# ============================================================================
# ОТЧЁТ
# ============================================================================
def _hl(h):
    return "без роста" if h == 0 else ("среднее за всю историю" if np.isinf(h) else "полураспад %g мес." % h)


def describe(ch: dict) -> str:
    """Параметры выбранного варианта по-русски."""
    a, tau = ch["nc"]
    corr = "без поправки по факту" if a == 0 else f"поправка по факту: α_max {a:g}, τ {tau:g}"
    m = ch["method"]
    if m == "mix":
        parts = " + ".join(f"{w:g}·{NAMES[t]}" for w, (t, _) in zip(ch["w"], ch["parts"]) if w > 0)
        return f"{parts}; {corr}"
    if m in ("cal", "cal_d"):
        h, w0 = ch["cfg"]
        lvl = ("уровень не подстраивается" if np.isinf(w0) else
               f"уровень: {'среднее за всю историю' if np.isinf(h) else 'полураспад %g мес.' % h}, тяга к 1 — {w0:g}")
        return f"{lvl}; {corr}"
    if m == "trend":
        return f"тренд: {_hl(ch['cfg'])}; {corr}"
    return corr


def honest_table(h: pd.DataFrame, methods: list) -> pd.DataFrame:
    res = {}
    for name in methods + ["mix", "auto"]:
        g = h[h["method"] == name]
        if g.empty:
            continue
        by_k = g.groupby("known_days")["err_eff_pp"]
        mae = by_k.apply(lambda s: s.abs().mean())
        within = by_k.apply(lambda s: (s.abs() <= 1).mean())
        g0 = g[g["known_days"] == 0]
        res[NAMES[name]] = {
            **{f"конец мес., известно {k} дн.": f"{mae.get(k, np.nan):.2f} п.п." for k in REPORT_KNOWN_DAYS},
            "месяцев ≤ 1 п.п. (известно 0 / 5 / 10 / 20 дн.)":
                " / ".join(f"{100 * within.get(k, np.nan):.0f}%" for k in (0, 5, 10, 20)),
            "на даты, прогноз с 1-го: ср.|разница|": f"{g0['dates_pp'].mean():.2f} п.п.",
            "то же с 5-го числа": f"{g0['dates_pp_5'].mean():.2f} п.п.",
            "на даты, прогноз с 1-го: дат ≤ 1 п.п.": f"{100 * g0['dates_within1'].mean():.0f}%",
            "дневной WAPE": f"{100 * g0['wape'].mean():.1f}%",
        }
    return pd.DataFrame(res)


def _reach(h: pd.DataFrame, share=0.8):
    """С какого числа известных дней разница эффективности на конец месяца ≤ 1 п.п. в доле share месяцев («Авто»)."""
    a = h[h["method"] == "auto"]
    within = a.groupby("known_days")["err_eff_pp"].apply(lambda s: (s.abs() <= 1).mean())
    ok = within[within >= share]
    return int(ok.index.min()) if len(ok) else None


def _curve(h: pd.DataFrame, name: str) -> pd.Series:
    return h[h["method"] == name].groupby("known_days")["err_eff_pp"].apply(lambda s: s.abs().mean())


# ============================================================================
# ВЫГРУЗКА В EXCEL: ТРЕУГОЛЬНИКИ ЭФФЕКТИВНОСТИ
# ============================================================================
XL_PCT, XL_PP, XL_MONEY, XL_DATE = "0.0%", "+0.00;-0.00;0.00", "#,##0", "DD.MM.YYYY"
XL_DAILY = {"month": "месяц", "date": "дата", "is_workday": "рабочий день", "has_fact": "факт есть", "plan": "план",
            "plan_is_forecast": "план — прогноз", "fact": "факт", "forecast": "прогноз без факта месяца",
            "fact_plus_forecast": "факт + прогноз остатка", "cum_plan": "план с начала месяца",
            "cum_fact": "факт с начала месяца", "cum_forecast": "прогноз с начала месяца",
            "cum_fact_plus_forecast": "факт + прогноз с начала месяца", "eff_fact_pct": "эфф. факт, %",
            "eff_forecast_pct": "эфф. прогноз, %", "eff_fact_plus_forecast_pct": "эфф. факт + прогноз, %",
            "eff_diff_pp": "прогноз − факт, п.п.", "forecast_err_pct": "ошибка прогноза суммы, %"}


def _day_table(dates, values) -> pd.DataFrame:
    """Строки — месяцы, столбцы — числа месяца 1..31."""
    dt = pd.DatetimeIndex(pd.to_datetime(dates))
    t = pd.DataFrame({"ym": dt.to_period("M").astype(str), "day": dt.day, "v": np.asarray(values, dtype=float)})
    return t.pivot(index="ym", columns="day", values="v").reindex(columns=range(1, 32))


def _xl(v):
    """Значение для ячейки Excel."""
    if isinstance(v, (list, tuple, np.ndarray)):
        return str(v)
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return None
    if isinstance(v, (pd.Timestamp, np.datetime64)):
        return pd.Timestamp(v).to_pydatetime()
    if isinstance(v, (bool, np.bool_)):
        return bool(v)
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, (float, np.floating)):
        return float(v) if np.isfinite(v) else None
    return v if isinstance(v, (str, int)) else str(v)


def _put_table(ws, df: pd.DataFrame, r0: int, title=None, note=None, fmt=None, style=None, index=True) -> int:
    """Таблица на лист с строки r0: заголовок, пояснение, шапка, строки. fmt(столбец) — формат чисел столбца,
    style(строка, столбец, значение) — (заливка, шрифт, формат) ячейки. Возвращает следующую свободную строку."""
    from openpyxl.styles import Alignment, Font, PatternFill
    r = r0
    if title:
        ws.cell(r, 1, title).font = Font(bold=True, size=12)
        r += 1
    if note:
        ws.cell(r, 1, note).font = Font(italic=True, color="52514E")
        r += 1
    head_font, head_fill = Font(bold=True), PatternFill("solid", fgColor="E9E9E6")
    head_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    names = ([df.index.name or ""] if index else []) + [str(c) for c in df.columns]
    for j, name in enumerate(names, start=1):
        c = ws.cell(r, j, name)
        c.font, c.fill, c.alignment = head_font, head_fill, head_align
    c0 = 2 if index else 1
    for i, (idx, row) in enumerate(df.iterrows(), start=r + 1):
        if index:
            ws.cell(i, 1, _xl(idx)).font = Font(bold=True)
        for j, col in enumerate(df.columns, start=c0):
            c = ws.cell(i, j, _xl(row[col]))
            f = fmt(col) if fmt else None
            if isinstance(c.value, datetime.datetime):
                f = XL_DATE
            if style:
                fill, font, f2 = style(idx, col, row[col])
                if fill is not None:
                    c.fill = fill
                if font is not None:
                    c.font = font
                f = f2 or f
            if f:
                c.number_format = f
    return r + len(df) + 2


def _err_fill(st: dict, e):
    """Заливка по |разнице| в п.п.: ≤ 1 — зелёная, ≤ 2 — жёлтая, больше — красная."""
    if e is None or not np.isfinite(e):
        return None
    return st["ok"] if abs(e) <= 1 else (st["mid"] if abs(e) <= 2 else st["bad"])


def export_excel(path: str, daily: pd.DataFrame, fc: pd.DataFrame, fsum: pd.DataFrame, paths: dict, bt: dict,
                 dev: dict, tables: list):
    """Excel в формате треугольника: строки — месяцы, столбцы — числа месяца 1..31, в ячейке — эффективность на эту
    дату (факт с начала месяца / план с начала месяца).
      daily  — дата, план, факт по всем месяцам (факт пустой, где его ещё нет);
      fc     — прогноз по дням на прогнозные месяцы (production), fsum — сводка по ним;
      paths  — честный бэктест «Авто»: {месяц: накопленный прогноз [k известных дней, день]};
      bt     — {месяц: (план по дням, факт по дням, даты)} для месяцев бэктеста;
      dev    — {месяц: прогноз итога месяца в деньгах при k = 0, 1, … известных днях} для прогнозных месяцев;
      tables — [(лист, [(заголовок, таблица), …]), …] — сводки."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("Excel не записан: нужен пакет openpyxl (pip install openpyxl).")
        return
    st = dict(fc_fill=PatternFill("solid", fgColor="DCEBFA"), fc_font=Font(color="1F5FAD", italic=True),
              none=PatternFill("solid", fgColor="F2F2F2"), ok=PatternFill("solid", fgColor="D9F2E3"),
              mid=PatternFill("solid", fgColor="FFF2CC"), bad=PatternFill("solid", fgColor="F8D7D3"))
    days = list(range(1, 32))

    def dim(m):
        return pd.Period(m, "M").days_in_month if m[:1].isdigit() else 31

    def is_day(col):
        return isinstance(col, (int, np.integer))

    fcm = {m: o for m, o in fc.groupby("month", sort=False)}
    bt_keys = [str(m) for m in sorted(paths)]
    btd = {str(m): v for m, v in bt.items()}
    wb = Workbook()

    # 1. Эффективность на каждую дату: факт, где его ещё нет — прогноз
    d = daily.copy()
    d["ym"] = d["date"].dt.to_period("M").astype(str)
    g = d.groupby("ym")
    with np.errstate(divide="ignore", invalid="ignore"):
        eff = (g["fact"].cumsum() / g["plan"].cumsum().replace(0, np.nan)).to_numpy()
    E = _day_table(d["date"], eff)
    F = pd.DataFrame(False, index=E.index, columns=days)
    status, end, plan_m = {}, {}, g["plan"].sum()
    full = g["fact"].count() == g.size()
    for m in E.index:
        if m not in fcm and full.get(m, False):
            status[m], end[m] = "факт", g["fact"].sum()[m] / plan_m[m]
    for m, o in fcm.items():
        k = int(o["has_fact"].sum())
        E.loc[m] = _day_table(o["date"], o["eff_fact_plus_forecast_pct"] / 100).loc[m]
        F.loc[m] = _day_table(o["date"], ~o["has_fact"]).loc[m].fillna(0).astype(bool)
        status[m] = ("факт" if k == len(o) else "прогноз" if k == 0
                     else f"факт по {pd.Timestamp(o['date'].iloc[k - 1]):%d.%m}, дальше прогноз")
        status[m] += " (план — прогноз)" if o["plan_is_forecast"].any() else ""
        end[m] = o["fact_plus_forecast"].sum() / o["plan"].sum()
        plan_m[m] = o["plan"].sum()
    E = E.sort_index()
    F = F.reindex(E.index, fill_value=False)
    interval = dict(zip(fsum["месяц"].astype(str), fsum["80% интервал"])) if len(fsum) else {}
    T1 = pd.concat([pd.Series(status, name="Статус").reindex(E.index), E,
                    pd.Series(end, name="На конец месяца").reindex(E.index),
                    pd.Series(interval, name="80% интервал").reindex(E.index),
                    plan_m.rename("План месяца").reindex(E.index)], axis=1)
    T1.index.name = "Месяц"

    def s1(idx, col, v):
        if is_day(col):
            if col > dim(idx):
                return st["none"], None, None
            if F.loc[idx, col]:
                return st["fc_fill"], st["fc_font"], None
        elif col == "На конец месяца" and idx in fcm and F.loc[idx].any():
            return st["fc_fill"], st["fc_font"], None
        return None, None, None

    ws = wb.active
    ws.title = "Эффективность"
    _put_table(ws, T1, 1, title="Эффективность на дату: факт с начала месяца / план с начала месяца",
               note="Чёрным — факт, синим на голубом — прогноз (факт по сегодня + прогноз остатка с поправкой). "
                    "Столбцы 1–31 — число месяца.",
               fmt=lambda c: XL_PCT if is_day(c) or c == "На конец месяца" else (XL_MONEY if c == "План месяца" else None),
               style=s1)

    # 2. Прогноз, сделанный до начала месяца (без его факта)
    rows, ends = [], {}
    for m in bt_keys:
        plan, fact, dates = btd[m]
        f0 = paths[pd.Period(m, "M")][0]
        with np.errstate(divide="ignore", invalid="ignore"):
            rows.append(_day_table(dates, f0 / np.cumsum(plan)))
        ends[m] = (f0[-1] / plan.sum(), fact.sum() / plan.sum())
    for m, o in fcm.items():
        if m in ends:
            continue
        rows.append(_day_table(o["date"], o["eff_forecast_pct"] / 100))
        ends[m] = (o["forecast"].sum() / o["plan"].sum(),
                   o["fact"].sum() / o["plan"].sum() if o["has_fact"].all() else np.nan)
    if rows:
        P1 = pd.concat(rows)
        P1["Прогноз на конец"] = [ends[m][0] for m in P1.index]
        P1["Факт на конец"] = [ends[m][1] for m in P1.index]
        P1["Разница, п.п."] = 100 * (P1["Прогноз на конец"] - P1["Факт на конец"])
        P1.index.name = "Месяц"
        ws = wb.create_sheet("Прогноз с 1-го числа")
        _put_table(ws, P1, 1, title="Прогноз эффективности на дату, сделанный до начала месяца (без его факта)",
                   note="Прошлые месяцы — честный бэктест «Авто»: модели и выбор метода — только по предыдущим месяцам.",
                   fmt=lambda c: XL_PCT if is_day(c) or c in ("Прогноз на конец", "Факт на конец")
                   else (XL_PP if c == "Разница, п.п." else None),
                   style=lambda i, c, v: ((st["none"], None, None) if is_day(c) and c > dim(i) else
                                          (_err_fill(st, v), None, None) if c == "Разница, п.п." else (None, None, None)))

    # 3. Разница прогноза с 1-го числа и факта на каждую дату, п.п.
    rows = []
    for m in bt_keys:
        plan, fact, dates = btd[m]
        with np.errstate(divide="ignore", invalid="ignore"):
            rows.append(_day_table(dates, 100 * (paths[pd.Period(m, "M")][0] - np.cumsum(fact)) / np.cumsum(plan)))
    for m, o in fcm.items():
        if m not in bt_keys and o["has_fact"].any() and not o["has_fact"].all():
            rows.append(_day_table(o["date"], o["eff_diff_pp"].where(o["has_fact"])))
    if rows:
        D = pd.concat(rows)
        base = D.loc[[m for m in bt_keys if m in D.index]]
        D.loc["Средняя |разница|"] = base.abs().mean()
        D.loc["Месяцев ≤ 1 п.п."] = (base.abs() <= 1).sum() / base.notna().sum().replace(0, np.nan)
        D.index.name = "Месяц"

        def s3(idx, col, v):
            if idx == "Месяцев ≤ 1 п.п.":
                return None, Font(bold=True), XL_PCT
            if col > dim(idx):
                return st["none"], None, None
            return _err_fill(st, v), (Font(bold=True) if idx == "Средняя |разница|" else None), None

        ws = wb.create_sheet("Разница с фактом")
        _put_table(ws, D, 1, title="Прогноз с 1-го числа минус факт: разница эффективности на каждую дату, п.п.",
                   note="Зелёный — в пределах 1 п.п., жёлтый — до 2 п.п., красный — больше. Внизу — по дню месяца: "
                        "средняя |разница| и доля месяцев в пределах 1 п.п.",
                   fmt=lambda c: XL_PP, style=s3)

    # 4. Прогноз итога месяца по мере поступления факта: k — сколько дней месяца уже известно
    rows, facts = {}, {}
    for m in bt_keys:
        plan, fact, dates = btd[m]
        rows[m] = pd.Series(paths[pd.Period(m, "M")][:, -1] / plan.sum())
        facts[m] = fact.sum() / plan.sum()
    for T, v in dev.items():
        m = str(T)
        if m in rows or m not in fcm:
            continue
        o = fcm[m]
        rows[m] = pd.Series(np.asarray(v, dtype=float) / o["plan"].sum())
        facts[m] = o["fact"].sum() / o["plan"].sum() if o["has_fact"].all() else np.nan
    if rows:
        K = pd.DataFrame(rows).T.reindex(columns=range(31))
        K["Факт на конец"] = [facts[m] for m in K.index]
        K.index.name = "Месяц"

        def s4(idx, col, v):
            if col == "Факт на конец":
                return None, Font(bold=True), None
            if col >= dim(idx):
                return st["none"], None, None
            fact = K.loc[idx, "Факт на конец"]
            return _err_fill(st, 100 * (v - fact) if np.isfinite(fact) and v is not None else None), None, None

        ws = wb.create_sheet("Итог месяца по мере факта")
        _put_table(ws, K, 1, title="Прогноз эффективности на конец месяца, когда известно k дней месяца",
                   note="Столбец — сколько дней месяца уже известно (0 — прогноз до начала месяца). Цвет — разница "
                        "с фактом на конец месяца: зелёный ≤ 1 п.п., жёлтый ≤ 2 п.п., красный больше.",
                   fmt=lambda c: XL_PCT, style=s4)

    # сводки
    for name, parts in tables:
        ws = wb.create_sheet("".join(" " if ch in ":\\/?*[]" else ch for ch in name)[:31])
        r = 1
        for title, t in parts:
            daily_sheet = "date" in t.columns
            if daily_sheet:
                t = t.rename(columns=XL_DAILY)
            r = _put_table(ws, t, r, title=title, index=not daily_sheet,
                           fmt=lambda c: "0.00" if any(s in str(c) for s in ("%", "п.п.")) else
                           (XL_MONEY if str(c) in ("план", "факт", "прогноз без факта месяца", "факт + прогноз остатка")
                            or "с начала месяца" in str(c) else None))

    for ws in wb.worksheets:
        ws.column_dimensions["A"].width = 12
        for j in range(2, ws.max_column + 1):
            head = ws.cell(3 if ws.title in ("Эффективность", "Прогноз с 1-го числа", "Разница с фактом",
                                             "Итог месяца по мере факта") else 2, j).value
            ws.column_dimensions[get_column_letter(j)].width = (7.5 if str(head).isdigit() else 14)
    sheet = wb["Эффективность"]
    sheet.column_dimensions["B"].width = 30
    sheet.freeze_panes = "C4"
    for name in ("Прогноз с 1-го числа", "Разница с фактом", "Итог месяца по мере факта"):
        if name in wb.sheetnames:
            wb[name].freeze_panes = "B4"
    wb.save(path)
    print(f"Excel: {path}")


# ============================================================================
# ЗАПУСК
# ============================================================================
def run(df: pd.DataFrame, last_fact_date=LAST_FACT_DATE, make_plots=True, prefix="", label=None, excel=None):
    d, last_fact_date, cur, targets = prepare(df, last_fact_date)
    has_debt = "debt" in d
    tag = f" [{label}]" if label else ""
    print(f"Данные{tag}: {d['date'].min():%d.%m.%Y} – {last_fact_date:%d.%m.%Y} (факт); "
          f"план загружен по {full_months(d, 'plan').index.max()}"
          + (f"; «к оплате» известно по {d.loc[d['debt_actual'].notna(), 'date'].max():%d.%m.%Y}" if has_debt else ""))

    # прогноз плана (и «к оплате») там, где их ещё нет
    plan_cfg = tune_series(d, "plan")
    d["plan"] = fill_series(d, "plan", plan_cfg)
    print(f"Прогноз плана, когда он не загружен: последний месяц ({plan_cfg['unit']}) × темп роста ({_hl(plan_cfg['h'])}), "
          f"раскладка {plan_cfg['profile']}. На прошлых месяцах ошибка итога плана {100 * plan_cfg['total_err']:.2f}%.")
    debt_cfg = None
    if has_debt and DEBT_KNOWN_AHEAD:
        short = [str(T) for T in targets if T >= cur and d.loc[d["ym"] == T, "debt_actual"].isna().any()]
        if short:
            warnings.warn(f"«К оплате» на {', '.join(short)} выгружено не на все дни: недостающие дни спрогнозированы "
                          f"по прошлым месяцам, и прогноз будет хуже бэктеста. Выгрузите «к оплате» на весь месяц "
                          f"или поставьте DEBT_KNOWN_AHEAD = False.")
    if has_debt:
        debt_cfg = tune_series(d, "debt")
        d["debt"] = fill_series(d, "debt", debt_cfg)
        if not DEBT_KNOWN_AHEAD:
            print(f"Прогноз «к оплате» на будущие дни: на прошлых месяцах ошибка итога {100 * debt_cfg['total_err']:.2f}%.")
    d = add_dist_features(d, "plan")
    if has_debt:
        d = add_dist_features(d, "debt")
    mt = month_table(d, last_fact_date)
    d["y"] = np.where(d["ym"].map(mt["plan_known"]).astype(bool), d["fact"] / d["plan_m"].replace(0, np.nan), np.nan)
    if has_debt:
        d["y_debt"] = np.where(d["ym"].map(mt["debt_known"]).astype(bool),
                               d["fact"] / d["debt_m"].replace(0, np.nan), np.nan)

    print("Бэктест...")
    months = backtest(d, mt, plan_cfg, debt_cfg)
    hist = histories(months)
    methods = method_list(daily_debt(d))
    sel, hon = {}, {}
    for sfx in SCENARIOS:
        SC, PR = score_table(months, hist, methods, sfx)
        sel[sfx] = select(SC, PR, months, list(range(len(months))), methods, sfx)
        hon[sfx] = honest(SC, PR, months, hist, methods, sfx)

    h0 = hon[""]["df"]
    hm = sorted(h0["ym"].unique())
    print(f"\nЧЕСТНЫЙ БЭКТЕСТ{tag}: каждый месяц ({len(hm)} мес., {hm[0]} – {hm[-1]}) прогнозируется методом и "
          f"параметрами, выбранными только по предыдущим месяцам. Разница эффективности прогноза и факта, п.п. "
          f"(эффективность = факт / план; прогноз / план).")
    for sfx, lab in SCENARIOS.items():
        print(f"\nЧестный бэктест, {lab}:")
        print(honest_table(hon[sfx]["df"], methods).to_string())
        r = _reach(hon[sfx]["df"])
        print("«Авто»: ≤ 1 п.п. на конец месяца в 80% месяцев — "
              + (f"когда известно {r} дн. месяца." if r is not None else "не достигается."))

    for sfx, lab in SCENARIOS.items():
        best, winner = sel[sfx]
        print(f"\nРабочий прогноз, {lab}: параметры подобраны по всем {len(months)} мес. бэктеста "
              f"(ошибка на подборе оптимистична — ориентируйтесь на честный бэктест выше).")
        t = pd.DataFrame({NAMES[m]: {"параметры": describe(b), "ошибка на подборе": f"{b['score']:.2f} п.п."}
                          for m, b in best.items()}).T
        print(t.to_string())
        print(f"→ в прогнозе: {NAMES[winner]}")

    a = h0[h0["method"] == "auto"]
    by_month = (a[a["known_days"].isin(REPORT_KNOWN_DAYS)]
                .pivot(index="ym", columns="known_days", values="err_eff_pp").round(2))
    by_month.insert(0, "эфф. факт, %", a.groupby("ym")["eff_fact"].first().round(1))
    by_month.insert(0, "метод", pd.Series(hon[""]["picks"]).rename(index=str).map(NAMES))
    pe = pd.Series({str(M["ym"]): 100 * (M["P_pu"] / M["P"] - 1) for M in months})
    by_month["ошибка прогноза плана, %"] = pe.round(2)
    print("\n«Авто» по месяцам, план загружен: разница эффективности на конец месяца, п.п. (столбцы — сколько дней "
          "известно; + значит прогноз выше факта). «метод» — что выбрано по предыдущим месяцам. Последний столбец — "
          "насколько ошибся бы прогноз плана, если бы KPI ещё не был загружен:")
    print(by_month.to_string())
    if daily_debt(d):
        rel = a.groupby("known_days")["err_eom"].apply(lambda s: 100 * s.abs().mean())
        print("Recovery месяца («Авто», честно; «к оплате» известно): средняя |относительная ошибка| — "
              + ", ".join(f"{k} дн.: {rel.get(k, np.nan):.2f}%" for k in (0, 5, 10, 20)))

    fc, fsum, models, dev = production(d, mt, cur, targets, sel, hist, hon)
    print(f"\nПрогноз{tag}:")
    print(fsum.to_string(index=False))

    winner = sel[""][1]
    main = sel[""][0][winner]
    key = (main["parts"][0][0] if main["method"] == "mix" else main["method"])
    key = {"boost": "boost", "cal": "cal", "trend": "cal", "boost_d": "boost_d", "cal_d": "cal_d"}.get(key, "boost")
    ms = models.get(cur, models[max(models)])
    model = ms[key]
    if getattr(model, "feature_importances_", None) is not None:
        imp = pd.Series(model.feature_importances_, index=model.feats_).sort_values(ascending=False)
        print(f"\nВажность признаков модели «{NAMES.get(key, key)}» (топ-12):")
        print(imp.head(12).round(1).to_string())
    if "boost_d" in ms:
        kw = ms["boost_d"].kernel.weights()
        print("\nКак платят относительно даты «к оплате» (доля оплат, оценка по истории; + — за столько дней до даты, "
              "− — через столько дней после):")
        print("  " + "  ".join(f"{lag:+d}: {100 * v:.0f}%" for lag, v in kw.items() if v >= 0.005))

    for m, g in fc.groupby("month"):
        g.to_csv(f"{OUT_DIR}/forecast_{prefix}{m}.csv", index=False, encoding="utf-8-sig")
        if make_plots:
            plot_month(g, f"{OUT_DIR}/forecast_{prefix}{m}.png")
    pd.concat([h.assign(scenario=SCENARIOS[s]) for s, h in ((s, hon[s]["df"]) for s in SCENARIOS)]).to_csv(
        f"{OUT_DIR}/{prefix}backtest_errors.csv", index=False, encoding="utf-8-sig")
    if make_plots:
        curves = {"Авто, план загружен": (_curve(h0, "auto"), BLUE, "-"),
                  "Авто, план ещё не загружен": (_curve(hon["_pu"]["df"], "auto"), BLUE, (0, (4, 3))),
                  f"{NAMES['ref']} (текущий подход)": (_curve(h0, "ref"), ORANGE, "-")}
        if daily_debt(d):
            curves[NAMES["boost_d"]] = (_curve(h0, "boost_d"), AQUA, "-")
        plot_backtest(curves, f"{OUT_DIR}/{prefix}backtest_error_by_day.png")
    tables = [("Точность методов", [(f"Честный бэктест, {lab}: разница эффективности прогноза и факта, п.п.",
                                      honest_table(hon[s]["df"], methods)) for s, lab in SCENARIOS.items()]),
              ("Авто по месяцам", [("«Авто» по месяцам, план загружен: разница эффективности на конец месяца, п.п. "
                                    "(+ — прогноз выше факта)",
                                    by_month.rename_axis("Месяц").rename(
                                        columns={k: f"известно {k} дн., п.п." for k in REPORT_KNOWN_DAYS}))])]
    if EXCEL if excel is None else excel:
        export_excel(f"{OUT_DIR}/{prefix}effectiveness.xlsx", d.loc[d["ym"] <= max(targets), ["date", "plan", "fact"]],
                     fc, fsum, hon[""]["paths"], {M["ym"]: (M["plan"], M["fact"], M["date"]) for M in months}, dev,
                     [("Сводка прогноза", [("Прогноз эффективности на конец месяца", fsum.set_index("месяц"))])]
                     + tables + [("Прогноз по дням", [("Прогноз по дням", fc)])])
    return dict(daily=d, months=mt, backtest=months, hist=hist, selection=sel, honest=hon, winner=winner,
                plan_cfg=plan_cfg, debt_cfg=debt_cfg, by_month=by_month, forecast=fc, forecast_summary=fsum,
                models=models, dev=dev, tables=tables, cur=cur, targets=targets)


# ============================================================================
# ПРОДУКТЫ: КАЖДЫЙ И ИТОГ
# ============================================================================
def bottom_up_honest(results: list, sfx: str) -> dict:
    """Честный бэктест суммы прогнозов продуктов: пути «Авто» каждого продукта (уже с поправкой по факту продукта)
    складываются, эффективность — к сумме планов."""
    common = set.intersection(*[set(r["honest"][sfx]["paths"]) for r in results])
    rows, paths = [], {}
    for ym in sorted(common):
        Ms = [next(M for M in r["backtest"] if M["ym"] == ym) for r in results]
        cums = sum(r["honest"][sfx]["paths"][ym] for r in results)
        fact = sum(M["fact"] for M in Ms)
        for row in path_rows(fact, cums, sum(M["plan"] for M in Ms), sum(M["plan" + sfx] for M in Ms)):
            rows.append(dict(method="auto", ym=str(ym), **row))
        paths[ym] = cums
    return dict(df=pd.DataFrame(rows), paths=paths)


def run_segments(df: pd.DataFrame, segment_col=None, make_plots=True):
    """Прогноз эффективности по каждому продукту/сегменту и по итогу. Итог считается двумя способами — моделью на
    сумме продуктов и суммой прогнозов продуктов; в прогноз итога идёт тот, что точнее в честном бэктесте."""
    segment_col = segment_col or SEGMENT_COL
    if not segment_col:
        raise ValueError("Укажите колонку продукта: run_segments(df, 'product') или SEGMENT_COL")
    cols = [DATE_COL, FACT_COL, PLAN_COL] + ([DEBT_COL] if DEBT_COL else [])
    res, failed = {}, []
    for s in sorted(df[segment_col].dropna().unique(), key=str):
        print(f"\n{'=' * 110}\nПРОДУКТ: {s}\n{'=' * 110}")
        try:
            res[s] = run(df.loc[df[segment_col] == s, cols], make_plots=make_plots, prefix=f"{s}_", label=str(s))
        except ValueError as e:     # например, у нового продукта мало истории
            print(f"Продукт {s} пропущен: {e}")
            failed.append(str(s))
    print(f"\n{'=' * 110}\nИТОГ: модель на сумме продуктов\n{'=' * 110}")
    total = df[cols].groupby(DATE_COL, as_index=False).sum(min_count=1)
    tot = run(total, make_plots=make_plots, prefix="total_", label="итог", excel=False)

    segs = [res[s] for s in res]
    print(f"\n{'=' * 110}\nИТОГ: модель на сумме против суммы прогнозов продуктов (честный бэктест, «Авто»)\n{'=' * 110}")
    choice = {sfx: ("модель на сумме", None) for sfx in SCENARIOS}
    compare = []
    if failed or not segs:
        print(f"Сумма продуктов не считается: не хватает прогноза по {', '.join(failed) or 'продуктам'}. "
              f"Итог — модель на сумме.")
    for sfx, lab in (SCENARIOS.items() if segs and not failed else []):
        bu = bottom_up_honest(segs, sfx)
        td = tot["honest"][sfx]["df"]
        td = td[(td["method"] == "auto") & td["ym"].isin(bu["df"]["ym"].unique())]
        t = pd.DataFrame({"модель на сумме": _curve(td, "auto"), "сумма продуктов": _curve(bu["df"], "auto")})
        t = t.loc[[k for k in REPORT_KNOWN_DAYS if k in t.index]].round(2)
        t.index = [f"конец мес., известно {k} дн., п.п." for k in t.index]
        s_td, s_bu = td["err_eff_pp"].abs().mean(), bu["df"]["err_eff_pp"].abs().mean()
        choice[sfx] = ("сумма продуктов", bu) if s_bu < s_td else ("модель на сумме", None)
        print(f"\n{lab} (средняя по всем дням |разница|: модель на сумме {s_td:.2f} п.п., сумма продуктов {s_bu:.2f} п.п.):")
        print(t.to_string())
        print(f"→ для итога берём: {choice[sfx][0]}")
        compare.append((f"Итог, {lab}: модель на сумме против суммы продуктов (честный бэктест «Авто», п.п.) — "
                        f"берём: {choice[sfx][0]}", t))

    # прогноз итога: сумма прогнозов продуктов или модель на сумме — по сценарию месяца
    fcs = [res[s]["forecast"].assign(segment=str(s)) for s in res]
    allf = pd.concat(fcs, ignore_index=True)
    out, summary = [], []
    for m, g_tot in tot["forecast"].groupby("month", sort=False):
        sc = "_pu" if g_tot["plan_is_forecast"].any() else ""
        name, bu = choice[sc]
        if bu is None:
            o = g_tot.copy()
            h = tot["honest"][sc]["df"]
        else:
            g = allf[allf["month"] == m]
            o = (g.groupby("date", as_index=False)
                 .agg(month=("month", "first"), is_workday=("is_workday", "first"), has_fact=("has_fact", "all"),
                      plan=("plan", "sum"), plan_is_forecast=("plan_is_forecast", "any"),
                      fact=("fact", lambda s: s.sum(min_count=1)), forecast=("forecast", "sum"),
                      fact_plus_forecast=("fact_plus_forecast", "sum")))
            o = add_effs(o)
            h = bu["df"]
        out.append(o.assign(total_method=name))
        summary.append({**month_summary(o, h, bool(sc), name)})
    fc_total = pd.concat(out, ignore_index=True)
    fsum_total = pd.DataFrame(summary)
    print("\nПрогноз ИТОГА:")
    print(fsum_total.to_string(index=False))
    for m, g in fc_total.groupby("month"):
        g.to_csv(f"{OUT_DIR}/forecast_ИТОГ_{m}.csv", index=False, encoding="utf-8-sig")
        if make_plots:
            plot_month(g, f"{OUT_DIR}/forecast_ИТОГ_{m}.png")
    rows = [r["forecast_summary"].assign(продукт=str(s)) for s, r in res.items()]
    allsum = pd.concat(rows + [fsum_total.assign(продукт="ИТОГ")], ignore_index=True)
    allsum = allsum[["продукт"] + [c for c in allsum.columns if c != "продукт"]]
    allsum.to_csv(f"{OUT_DIR}/forecast_summary_by_product.csv", index=False, encoding="utf-8-sig")
    print("\nСводка по продуктам и итогу:")
    print(allsum[["продукт", "месяц", "дней факта", "эфф.: прогноз без факта", "эфф.: факт + прогноз остатка",
                  "80% интервал", "метод"]].to_string(index=False))
    if EXCEL:  # Excel итога: прогноз — выбранный способ; бэктест и сходимость — тем же способом, что и для плана загружен
        bu = choice[""][1]
        paths = bu["paths"] if bu is not None else tot["honest"][""]["paths"]
        dev = dict(tot["dev"])
        for T in dev:   # для месяца — тем же способом, что и прогноз итога (зависит от того, загружен ли план)
            sc = "_pu" if fc_total.loc[fc_total["month"] == str(T), "plan_is_forecast"].any() else ""
            if choice[sc][1] is not None and all(T in r["dev"] for r in segs):
                dev[T] = sum(np.asarray(r["dev"][T], dtype=float) for r in segs)
        export_excel(f"{OUT_DIR}/ИТОГ_effectiveness.xlsx",
                     tot["daily"].loc[tot["daily"]["ym"] <= max(tot["targets"]), ["date", "plan", "fact"]],
                     fc_total, fsum_total, paths,
                     {M["ym"]: (M["plan"], M["fact"], M["date"]) for M in tot["backtest"]}, dev,
                     [("Сводка по продуктам", [("Прогноз эффективности по продуктам и итогу", allsum.set_index("продукт"))]),
                      ("Итог — способ", compare or [("Итог — модель на сумме продуктов", pd.DataFrame())])]
                     + tot["tables"] + [("Прогноз по дням", [("Прогноз итога по дням", fc_total)])])
    return dict(segments=res, total_model=tot, total_forecast=fc_total, total_summary=fsum_total,
                total_choice={s: c[0] for s, c in choice.items()}, summary=allsum)


def load_data() -> pd.DataFrame:
    """Замените на свой запрос. Нужны колонки DATE_COL, FACT_COL, PLAN_COL (и DEBT_COL, SEGMENT_COL, если заданы)."""
    # import sqlalchemy as sa
    # engine = sa.create_engine("...")
    # return pd.read_sql("select pay_date as date, product, fact, plan, to_pay as debt from ...", engine)
    raise NotImplementedError("Заполните load_data() своим запросом к базе")


if __name__ == "__main__":
    data = load_data()
    res = run_segments(data) if SEGMENT_COL else run(data)
