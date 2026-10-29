"""
Прогноз RR1+ и выполнения плана СНИЗУ ВВЕРХ — по сегментам портфеля просрочки, на каждую дату месяца и на конец месяца.

ПОЧЕМУ ПО СЕГМЕНТАМ
  RR итога = Σ сбор сегментов / Σ w1, выполнение плана = Σ сбор / Σ target. Сбор сильно зависит от DPD-бакета, PVR-рейтинга
  и продукта: ранние бакеты за месяц собираются на десятки процентов, глубокие — на единицы. Состав портфеля на 1-е число
  известен (w1 по сегментам), поэтому сдвиг структуры — больше раннего или позднего долга — в прогнозе по итогу выглядит как
  непредсказуемый «уровень», а снизу вверх считается напрямую. profile() показывает, какая часть изменения RR от месяца
  к месяцу приходится на структуру.

ДАННЫЕ
  Сырые срезы RR1+ (…_data_raw): отчётный месяц, дата среза, сегмент (субпродукт × тербанк × DPD-бакет × PVR-бакет ×
  inflow), w1 — долг когорты на входе, w2 — остаток долга на дату среза, накопленные списания и цессии, target — план сбора
  на дату. Факт — строки is_forecast = 0; строки is_forecast = 1 — план (target, w1) на будущие даты.
  Сбор на дату = w1 − w2 − списания − цессии. RR = сбор / w1, выполнение плана = сбор / target.

КАК УСТРОЕН ПРОГНОЗ
  1. Срезы собираются в панель: месяц × сегмент (GROUP_COLS) × число месяца; в дни без среза — последний срез.
  2. Кривая сбора сегмента — доля w1 конца месяца, собранная к каждому дню, средневзвешенная по прошлым месяцам (вес падает
     с давностью), дни выровнены по календарю или по рабочим дням, с поправкой на сезон (тот же месяц год назад). Кривые
     маленьких сегментов подтягиваются к кривым родительских уровней (SHRINK_LEVELS) — чем меньше сегмент, тем сильнее.
  3. Ожидаемый сбор сегмента = кривая × w1. w1 портфеля на 1-е число известен; если w1 сегмента растёт внутри месяца
     (новые входы), его рост прогнозируется по прошлым месяцам.
  4. Внутри месяца ожидание поправляется по факту на сегодня — каждого сегмента или всего портфеля; доверие растёт с долей
     месяца, которая уже должна была прийти: α = α_max · c / (c + τ).
  5. Честный бэктест: для каждого месяца метод и параметры выбираются только по предыдущим месяцам. Методы:
       Референсные месяцы — текущий подход: средняя кривая выполнения плана 3 прошлых месяцев по числу месяца;
       Сверху вниз        — кривая сбора портфеля в целом × w1 (без структуры);
       Снизу вверх        — кривые сегментов × w1 сегментов;
       Авто               — каждый месяц лучший по предыдущим месяцам вариант.
     Выбор — по выполнению плана (оно строже: 1 п.п. выполнения ≈ 1% суммы сбора); RR показывается рядом.

ЗАПУСК (Jupyter)
  import rr_forecast as rr
  raw = rr.load_raw()          # или свой запрос: нужны колонки из COLS и GROUP_COLS
  rr.profile(raw)              # ПЕРВЫЙ ШАГ: как устроены данные — пришлите вывод
  res = rr.run(raw)            # бэктест, прогноз, Excel rr_effectiveness.xlsx
Нужны pandas, numpy, openpyxl (для Excel) и collections_forecast.py в той же папке (календарь и выгрузка в Excel).
"""

import itertools

import numpy as np
import pandas as pd

import collections_forecast as cf

# ============================================================================
# НАСТРОЙКИ
# ============================================================================
COLS = dict(month="report_month", date="report_dt", w1="w1", w2="w2", write_off="write_off", cession="cession",
            target="target", is_forecast="is_forecast")
GROUP_COLS = ["subproduct", "dpd_bucket", "pvr_bucket", "inflow"]   # сегмент прогноза; можно добавить "terbank"
SHRINK_LEVELS = [["subproduct", "dpd_bucket", "inflow"],             # родительские уровни, от мелкого к крупному;
                 ["dpd_bucket", "inflow"]]                            # над последним — портфель в целом
W1_VARY_TOL = 0.005            # w1 сегмента «растёт внутри месяца», если к концу месяца он обычно больше начального на 0,5%
BACKTEST_FROM = "2025-01"      # с какого месяца бэктест
HONEST_WARMUP = 6              # первые N месяцев бэктеста — только для выбора метода
HISTORY_MONTHS = 12            # сколько прошлых месяцев берётся в кривые
ALIGNS = ("calendar", "workday")              # выравнивание дней: по числу месяца / по доле прошедших рабочих дней
CURVE_HALFLIVES = (1, 3, 6, np.inf)           # вес прошлого месяца в кривой падает вдвое каждые N мес. (inf — поровну)
SHRINK_K = (0, 0.1, 1, 10)                    # сила подтягивания к родителю, в медианах размера сегмента (0 — нет)
SEASONAL = (0, 0.5, 1)                        # поправка на сезон: доля учёта того же месяца год назад
NOWCAST_MAX = (0.25, 0.5, 0.75, 1.0)          # сила поправки по факту месяца
NOWCAST_TAU = (0, 0.02, 0.05, 0.1, 0.2, 0.5, 1.0)
RATIO_CLIP = (0.5, 2.0)        # факт / ожидание на сегодня не выходит за эти пределы (в праздники ожидание почти нулевое)
REPORT_KNOWN_DAYS = (0, 1, 3, 5, 10, 15, 20, 25)
OUT_DIR = "."
EXCEL = True


def load_raw() -> pd.DataFrame:
    """Замените на свой запрос к таблице срезов RR1+. Достаточно колонок COLS и GROUP_COLS."""
    # import sqlite3
    # con = sqlite3.connect(r"путь\к\базе.db")
    # cols = ", ".join(list(COLS.values()) + sorted(set(GROUP_COLS) | {"terbank", "dt_in"}))
    # return pd.read_sql(f"select {cols} from <таблица>_data_raw", con)
    raise NotImplementedError("Заполните load_raw() своим запросом")


# ============================================================================
# ДАННЫЕ → ПАНЕЛЬ: месяц × сегмент × число месяца
# ============================================================================
def _prep(raw: pd.DataFrame) -> pd.DataFrame:
    c = COLS
    need = [c[k] for k in ("month", "date", "w1", "w2", "write_off", "cession", "target")] + GROUP_COLS
    miss = [x for x in need if x not in raw.columns]
    if miss:
        raise KeyError(f"Нет колонок: {miss}. Поправьте COLS / GROUP_COLS")
    r = pd.DataFrame({"rm": pd.to_datetime(raw[c["month"]]).dt.to_period("M"),
                      "dt": pd.to_datetime(raw[c["date"]]).dt.normalize()})
    for k in ("w1", "w2", "write_off", "cession", "target"):
        r[k] = pd.to_numeric(raw[c[k]], errors="coerce").fillna(0.0).to_numpy()
    fc = c.get("is_forecast")
    r["fc"] = pd.to_numeric(raw[fc], errors="coerce").fillna(0).to_numpy() if fc in raw else 0
    for g in GROUP_COLS:
        r[g] = raw[g].astype(str).to_numpy()
    off = (r["dt"] - r["rm"].dt.start_time).dt.days
    r["day"] = np.clip(off + 1, 1, r["rm"].dt.days_in_month)   # срез после конца месяца — это конец месяца
    return r


def _cube(df: pd.DataFrame, cols, mi: dict, M: int, G: int):
    """Суммы по месяц × сегмент × день: {колонка: [M, G, 31]} (NaN — нет строки) и признак среза [M, 31]."""
    a = df.groupby(["rm", "day", "g"], sort=False)[cols].sum()
    im = a.index.get_level_values(0).map(mi).to_numpy()
    idd = a.index.get_level_values(1).to_numpy().astype(int) - 1
    ig = a.index.get_level_values(2).to_numpy()
    out = {}
    for col in cols:
        x = np.full((M, G, 31), np.nan)
        x[im, ig, idd] = a[col].to_numpy()
        out[col] = x
    has = np.zeros((M, 31), bool)
    has[im, idd] = True
    return out, has


def _ffill_days(x: np.ndarray, upto: int) -> np.ndarray:
    """Последний срез на каждый день до upto (включительно), дальше — NaN."""
    f = pd.DataFrame(x[:, :upto + 1]).ffill(axis=1).to_numpy()
    out = np.full_like(x, np.nan)
    out[:, :upto + 1] = f
    return out


def build_panel(raw: pd.DataFrame) -> dict:
    r = _prep(raw)
    grp = r.groupby(GROUP_COLS, sort=True)
    r["g"] = grp.ngroup().to_numpy()
    groups = grp.size().reset_index()[GROUP_COLS]
    fact, fut = r[r["fc"] == 0], r[r["fc"] == 1]
    last_m = max(fact["rm"].max(), fut["rm"].max()) if len(fut) else fact["rm"].max()
    months = pd.period_range(fact["rm"].min(), last_m, freq="M")
    M, G = len(months), len(groups)
    mi = {m: i for i, m in enumerate(months)}
    cols = ["w1", "w2", "write_off", "cession", "target"]
    cube, has = _cube(fact, cols, mi, M, G)
    fcube, fhas = _cube(fut, ["w1", "target"], mi, M, G) if len(fut) else ({}, np.zeros((M, 31), bool))

    dim = np.array([m.days_in_month for m in months])
    last_day = np.array([np.flatnonzero(has[i]).max() if has[i].any() else -1 for i in range(M)])
    first_day = np.array([np.flatnonzero(has[i]).min() if has[i].any() else -1 for i in range(M)])
    last_fact_m = int(np.flatnonzero(has.any(axis=1)).max())
    complete = np.array([i < last_fact_m or (i == last_fact_m and last_day[i] >= dim[i] - 1) for i in range(M)])
    for i in range(M):
        if last_day[i] < 0:
            continue
        upto = dim[i] - 1 if complete[i] else last_day[i]
        for col in cols:
            x = _ffill_days(cube[col][i], upto)
            if col == "w1":   # до первого среза — w1 первого среза
                x[:, :first_day[i]] = x[:, [first_day[i]]]
            x[:, :upto + 1] = np.nan_to_num(x[:, :upto + 1])
            cube[col][i] = x
    for col in fcube:     # план на будущие даты: последнее известное значение
        for i in range(M):
            if fhas[i].any():
                x = pd.DataFrame(fcube[col][i]).ffill(axis=1).bfill(axis=1).to_numpy(copy=True)
                x[:, dim[i]:] = np.nan
                fcube[col][i] = np.nan_to_num(x)

    W1, TG = cube["w1"], cube["target"]
    C = W1 - cube["w2"] - cube["write_off"] - cube["cession"]
    cal = cf.build_calendar(months[0].start_time, months[-1].end_time.normalize()).set_index("date")
    wdnum, wdtot = np.zeros((M, 31)), np.zeros(M)
    for i, m in enumerate(months):
        cm = cal.loc[m.start_time:m.end_time.normalize()]
        wdnum[i, :dim[i]] = cm["wd_num"].to_numpy()
        wdtot[i] = cm["wd_total"].iloc[0]
    # w1 сегмента растёт внутри месяца? (медиана по полным месяцам отношения w1 конца к w1 первого среза)
    ratios = []
    for i in np.flatnonzero(complete):
        a, b = W1[i, :, first_day[i]], W1[i, :, dim[i] - 1]
        ratios.append(np.where(a > 0, b / np.where(a > 0, a, 1), np.where(b > 0, np.inf, 1.0)))
    varying = np.nanmedian(np.array(ratios), axis=0) > 1 + W1_VARY_TOL if ratios else np.zeros(G, bool)
    P = dict(months=months, groups=groups, W1=W1, C=C, TG=TG, has=has, dim=dim, first=first_day, last=last_day,
             complete=complete, wdnum=wdnum, wdtot=wdtot, varying=varying, fut=fcube, fhas=fhas)
    P["levels"] = _levels(P)
    tot_c = np.array([C[i, :, dim[i] - 1].sum() if complete[i] else np.nan for i in range(M)])
    tot_w = np.array([W1[i, :, dim[i] - 1].sum() if complete[i] else np.nan for i in range(M)])
    with np.errstate(divide="ignore", invalid="ignore"):
        P["rr_end"] = tot_c / tot_w
    return P


def _levels(P: dict) -> list:
    """Уровни группировки от сегмента к портфелю: коды, число групп, кубы C и W1, отображение в родителя."""
    g = P["groups"]
    specs = [list(GROUP_COLS)] + [list(s) for s in SHRINK_LEVELS] + [[]]
    out = []
    for cols in specs:
        codes = (g.groupby(cols, sort=True).ngroup().to_numpy() if cols else np.zeros(len(g), int))
        L = int(codes.max()) + 1
        Cl, Wl, Tl = (np.zeros((len(P["months"]), L, 31)) for _ in range(3))
        for i in range(len(P["months"])):
            np.add.at(Cl[i], codes, np.nan_to_num(P["C"][i]))
            np.add.at(Wl[i], codes, np.nan_to_num(P["W1"][i]))
            np.add.at(Tl[i], codes, np.nan_to_num(P["TG"][i]))
        out.append(dict(cols=cols, codes=codes, L=L, C=Cl, W=Wl, T=Tl))
    for a, b in zip(out[:-1], out[1:]):   # родитель каждой группы уровня a — группа уровня b
        parent = np.zeros(a["L"], int)
        parent[a["codes"]] = b["codes"]
        a["parent"] = parent
    return out


# ============================================================================
# КРИВЫЕ СБОРА
# ============================================================================
def _align(P: dict, i: int, j: int, how: str):
    """Для каждого дня месяца i — точки месяца j с тем же положением: индексы (со сдвигом +1; 0 — «до начала»,
    значение 0) и доля интерполяции между ними."""
    di, dj = P["dim"][i], P["dim"][j]
    if how == "calendar":
        lo = np.minimum(np.arange(di), dj - 1) + 1
        return lo, lo, np.zeros(di)
    wi, wj = P["wdnum"][i, :di], P["wdnum"][j, :dj]
    pos = wi / P["wdtot"][i]
    work = np.flatnonzero(np.diff(np.r_[0, wj]) > 0)              # рабочие дни месяца j
    x = np.r_[0.0, wj[work] / P["wdtot"][j]]
    idx = np.r_[work[0], work + 1]      # точка 0 — последний нерабочий день перед первым рабочим (или 0, если его нет)
    k = np.clip(np.searchsorted(x, pos, side="right") - 1, 0, len(x) - 1)
    k2 = np.minimum(k + 1, len(x) - 1)
    span = x[k2] - x[k]
    frac = np.where(span > 0, (pos - x[k]) / np.where(span > 0, span, 1), 0.0)
    return idx[k], idx[k2], frac


def _gather(v: np.ndarray, lo, hi, frac) -> np.ndarray:
    v0 = np.concatenate([np.zeros((v.shape[0], 1)), v], axis=1)
    return v0[:, lo] * (1 - frac) + v0[:, hi] * frac


def _curves(P: dict, i: int, how: str, h: float, cache: dict) -> list:
    """Для каждого уровня: rel [L, n] — доля w1 конца месяца, собранная к каждому дню месяца i (средневзвешенно по прошлым
    месяцам), wbar [L] — средний w1 конца месяца, wsh [L, n] — доля w1 конца месяца, набранная к дню, relp [L, n] —
    сбор к дню / план на конец месяца."""
    key = (i, how, h)
    if key in cache:
        return cache[key]
    n = P["dim"][i]
    past = [j for j in range(max(0, i - HISTORY_MONTHS), i) if P["complete"][j]]
    res = []
    for lev in P["levels"]:
        num, numw, den, dent, sw = (np.zeros((lev["L"], n)), np.zeros((lev["L"], n)), np.zeros(lev["L"]),
                                    np.zeros(lev["L"]), 0.0)
        for j in past:
            w = 1.0 if np.isinf(h) else 0.5 ** ((i - j - 1) / h)
            dj = P["dim"][j]
            lo, hi, fr = _align(P, i, j, how)
            num += w * _gather(lev["C"][j, :, :dj], lo, hi, fr)
            numw += w * _gather(lev["W"][j, :, :dj], lo, hi, fr)
            den += w * lev["W"][j, :, dj - 1]
            dent += w * lev["T"][j, :, dj - 1]
            sw += w
        with np.errstate(divide="ignore", invalid="ignore"):
            rel = np.where(den[:, None] > 0, num / den[:, None], np.nan)
            wsh = np.where(den[:, None] > 0, numw / den[:, None], np.nan)
            relp = np.where(dent[:, None] > 0, num / dent[:, None], np.nan)
        res.append((rel, den / max(sw, 1e-12), wsh, relp))
    cache[key] = res
    return res


def _shrunk(P: dict, curves: list, K: float, which: int = 0) -> np.ndarray:
    """Кривые сегментов (which = 0 — сбор / w1, 3 — сбор / план месяца), подтянутые к родителям:
    z = w / (w + K · медиана размера сегмента)."""
    w0 = curves[0][1]
    k_abs = K * (np.median(w0[w0 > 0]) if (w0 > 0).any() else 0.0)
    cur = np.nan_to_num(curves[-1][which])
    for lev, cv in zip(P["levels"][-2::-1], curves[-2::-1]):
        rel, wbar = cv[which], cv[1]
        par = cur[lev["parent"]]
        z = np.where(np.isfinite(rel).all(axis=1) & (wbar > 0), wbar / (wbar + k_abs) if k_abs > 0 else 1.0, 0.0)
        cur = z[:, None] * np.nan_to_num(rel) + (1 - z[:, None]) * par
    return cur


def _seasonal(P: dict, i: int) -> float:
    """Сезонный множитель: RR того же месяца год назад к среднему RR трёх месяцев перед ним."""
    rr = P["rr_end"]
    if i - 15 < 0 or not np.isfinite(rr[[i - 12, i - 13, i - 14, i - 15]]).all():
        return 1.0
    return float(rr[i - 12] / rr[[i - 13, i - 14, i - 15]].mean())


def _w1_prior(P: dict, i: int, curves: list):
    """w1 конца месяца по сегментам (известный на начало месяца; для растущих внутри месяца — по прошлым месяцам)
    и путь w1 по дням [G, n]."""
    rel, wbar, wsh = curves[0][:3]
    f = max(P["first"][i], 0)
    start = P["W1"][i, :, f]
    vary = P["varying"]
    wend = np.where(vary, np.nan_to_num(wbar), np.nan_to_num(start))
    share = np.where(vary[:, None], np.nan_to_num(wsh, nan=1.0), 1.0)
    return wend, share * wend[:, None]


# ============================================================================
# МЕТОДЫ
# ============================================================================
NAMES = {"ref": "Референсные месяцы (текущий подход)", "td": "Сверху вниз (кривая портфеля)",
         "bu": "Снизу вверх от сбора (RR сегментов)", "bp": "Снизу вверх от плана (выполнение сегментов)",
         "mix": "Смесь двух «снизу вверх»", "auto": "Авто"}
NC = np.array([(0.0, 0.0)] + [(a, t) for a in NOWCAST_MAX for t in NOWCAST_TAU])
METHOD_CFGS = {
    "ref": [None],
    "td": list(itertools.product(ALIGNS, CURVE_HALFLIVES, [0], SEASONAL)),
    "bu": list(itertools.product(ALIGNS, CURVE_HALFLIVES, SHRINK_K, SEASONAL)),
    "bp": list(itertools.product(ALIGNS, CURVE_HALFLIVES, SHRINK_K, [0])),
}
NC_LEVELS = {"ref": ["total"], "td": ["total"], "bu": ["group", "total"], "bp": ["group", "total"],
             "mix": ["group", "total"]}
MIX_W = (0.25, 0.5, 0.75)      # доля «от сбора» в смеси


def plan_known(P: dict, i: int, cache: dict) -> np.ndarray:
    """План сегментов [G, n] на каждый день месяца i, каким он известен на начало месяца: у сегментов с растущим w1 —
    целевая ставка × прогнозный путь w1 (фактический рост w1 заранее неизвестен)."""
    key = ("plan", i)
    if key not in cache:
        n = P["dim"][i]
        tg, w = np.nan_to_num(P["TG"][i, :, :n]), np.nan_to_num(P["W1"][i, :, :n])
        _, wp = _w1_prior(P, i, _curves(P, i, "calendar", np.inf, cache))
        with np.errstate(divide="ignore", invalid="ignore"):
            rate = np.where(w > 0, tg / np.where(w > 0, w, 1), 0.0)
        cache[key] = np.where(P["varying"][:, None], rate * wp, tg)
    return cache[key]


def expected(P: dict, i: int, meth: str, cfg, cache: dict, tg_g: np.ndarray):
    """Ожидаемый накопленный сбор по дням месяца i без его факта: E [G или 1, n] и путь w1 [G или 1, n].
    tg_g — план сегментов по дням, известный на начало месяца."""
    n = P["dim"][i]
    if meth == "mix":
        w, cbu, cbp = cfg
        Eu, wp = expected_cached(P, i, "bu", cbu, cache, tg_g)
        Ep, _ = expected_cached(P, i, "bp", cbp, cache, tg_g)
        return w * Eu + (1 - w) * Ep, wp
    tg_tot = tg_g.sum(axis=0)
    if meth == "ref":       # средняя кривая выполнения плана 3 прошлых месяцев по числу месяца × план месяца
        past = [j for j in range(i - 1, -1, -1) if P["complete"][j]][:3]
        comp = []
        for j in past:
            dj = P["dim"][j]
            c, t = P["C"][j, :, :dj].sum(axis=0), P["TG"][j, :, :dj].sum(axis=0)
            with np.errstate(divide="ignore", invalid="ignore"):
                cj = np.where(t > 0, c / t, np.nan)
            comp.append(cj[np.minimum(np.arange(n), dj - 1)])
        comp = np.nan_to_num(np.nanmean(comp, axis=0)) if comp else np.zeros(n)
        _, wp = _w1_prior(P, i, _curves(P, i, "calendar", np.inf, cache))
        return (comp * tg_tot)[None, :], wp.sum(axis=0)[None, :]
    how, h, K, s = cfg
    cur = _curves(P, i, how, h, cache)
    sf = _seasonal(P, i) ** s
    wend, wp = _w1_prior(P, i, cur)
    if meth == "td":
        return (np.nan_to_num(cur[-1][0]) * sf * wend.sum())[:1], wp.sum(axis=0)[None, :]
    if meth == "bp":        # доля плана месяца, собранная к дню (по прошлым месяцам) × план сегмента на конец месяца
        return _shrunk(P, cur, K, which=3) * tg_g[:, -1:], wp
    return _shrunk(P, cur, K) * sf * wend[:, None], wp


def expected_cached(P: dict, i: int, meth: str, cfg, cache: dict, tg_g: np.ndarray):
    key = ("E", i, meth, cfg)
    if key not in cache:
        cache[key] = expected(P, i, meth, cfg, cache, tg_g)
    return cache[key]


def _eom(A: np.ndarray, E: np.ndarray, nc: np.ndarray) -> np.ndarray:
    """Прогноз итога месяца [конфигурация поправки, k] при k = 0..n-1 известных днях; A, E — [G, n]."""
    G, n = A.shape
    Aq = np.concatenate([np.zeros((G, 1)), A[:, :-1]], axis=1)
    Eq = np.concatenate([np.zeros((G, 1)), E[:, :-1]], axis=1)
    Ee = E[:, -1:]
    with np.errstate(divide="ignore", invalid="ignore"):
        ratio = np.clip(np.where(Eq > 0, Aq / np.where(Eq > 0, Eq, 1), 1.0), *RATIO_CLIP)
        c = np.where(Ee > 0, Eq / np.where(Ee > 0, Ee, 1), 0.0)
        w = np.where(c[None] > 0, c[None] / (c[None] + nc[:, 1, None, None]), 0.0)
    return (Aq[None] + (Ee - Eq)[None] * (1 + nc[:, 0, None, None] * w * (ratio[None] - 1))).sum(axis=1)


def _paths(A: np.ndarray, E: np.ndarray, nc) -> np.ndarray:
    """Накопленный сбор с поправкой [k, день] при k = 0..n-1 известных днях (сумма по группам)."""
    G, n = A.shape
    amax, tau = nc
    out = np.empty((n, n))
    Ee = E[:, -1]
    for k in range(n):
        if k == 0:
            out[0] = E.sum(axis=0)
            continue
        aq, eq = A[:, k - 1], E[:, k - 1]
        with np.errstate(divide="ignore", invalid="ignore"):
            ratio = np.clip(np.where(eq > 0, aq / np.where(eq > 0, eq, 1), 1.0), *RATIO_CLIP)
            c = np.where(Ee > 0, eq / np.where(Ee > 0, Ee, 1), 0.0)
            w = np.where(c > 0, c / (c + tau), 0.0)
        fc = aq[:, None] + (E[:, k:] - eq[:, None]) * (1 + amax * w * (ratio - 1))[:, None]
        out[k] = np.r_[A[:, :k].sum(axis=0), fc.sum(axis=0)]
    return out


def _w1_paths(W: np.ndarray, Wp: np.ndarray) -> np.ndarray:
    """w1 портфеля [k, день]: известный по сегодня; дальше — прогноз с поправкой на отклонение от ожидаемого роста."""
    G, n = W.shape
    out = np.empty((n, n))
    out[0] = Wp.sum(axis=0)
    for k in range(1, n):
        wq, pq = W[:, k - 1], Wp[:, k - 1]
        with np.errstate(divide="ignore", invalid="ignore"):
            r = np.clip(np.where(pq > 0, wq / np.where(pq > 0, pq, 1), 1.0), 0.5, 2)
        out[k] = np.r_[W[:, :k].sum(axis=0), (wq[:, None] + (Wp[:, k:] - pq[:, None]) * r[:, None]).sum(axis=0)]
    return out


def _actual(P: dict, i: int):
    n = P["dim"][i]
    return P["C"][i, :, :n], P["W1"][i, :, :n], P["TG"][i, :, :n].sum(axis=0)


def score_month(P: dict, i: int, cache: dict) -> dict:
    """Средняя по k |разница выполнения плана на конец месяца|, п.п.: {метод: [конфигурация, уровень × поправка]}."""
    tg_g = plan_known(P, i, cache)
    out = {}
    for meth, cfgs in METHOD_CFGS.items():
        out[meth] = np.array([_score(P, i, meth, expected(P, i, meth, cfg, cache, tg_g)[0]) for cfg in cfgs])
    return out


def _score(P: dict, i: int, meth: str, E: np.ndarray) -> np.ndarray:
    """Средняя по k |разница выполнения плана на конец месяца| для всех уровней и вариантов поправки по факту."""
    A, W, TG = _actual(P, i)
    a_end, t_end = A[:, -1].sum(), TG[-1]
    errs = []
    for lev in NC_LEVELS[meth]:
        nc = NC[:1] if meth == "ref" else NC
        eom = _eom(A, E, nc) if lev == "group" and E.shape[0] > 1 else _eom(A.sum(axis=0)[None], E.sum(axis=0)[None], nc)
        errs.append(np.abs(100 * (eom - a_end) / t_end).mean(axis=1))
    return np.concatenate(errs)


def _nc_list(meth: str) -> list:
    return [(lev, tuple(x)) for lev in NC_LEVELS[meth] for x in (NC[:1] if meth == "ref" else NC)]


def select(P: dict, test: list, scores: list, idx: list, cache: dict) -> tuple:
    """Лучшая конфигурация каждого метода, смесь двух «снизу вверх» и победитель — по месяцам test[idx]."""
    best = {}
    for meth in METHOD_CFGS:
        s = np.mean([scores[t][meth] for t in idx], axis=0)
        ci, ni = np.unravel_index(np.nanargmin(s), s.shape)
        lev, nc = _nc_list(meth)[ni]
        best[meth] = dict(method=meth, cfg=METHOD_CFGS[meth][ci], level=lev, nc=nc, score=float(s[ci, ni]))
    sc = []
    for w in MIX_W:
        cfg = (w, best["bu"]["cfg"], best["bp"]["cfg"])
        sc.append(np.mean([_score(P, test[t], "mix", expected(P, test[t], "mix", cfg, cache,
                                                              plan_known(P, test[t], cache))[0]) for t in idx], axis=0))
    sc = np.array(sc)
    wi, ni = np.unravel_index(np.nanargmin(sc), sc.shape)
    lev, nc = _nc_list("mix")[ni]
    best["mix"] = dict(method="mix", cfg=(MIX_W[wi], best["bu"]["cfg"], best["bp"]["cfg"]), level=lev, nc=nc,
                       score=float(sc[wi, ni]))
    winner = min(best, key=lambda m: best[m]["score"])
    return best, winner


def forecast_paths(P: dict, i: int, ch: dict, cache: dict, A: np.ndarray, W: np.ndarray, tg_g: np.ndarray):
    """Пути накопленного сбора и w1 [k, день] выбранного варианта (с поправкой по факту при k известных днях)."""
    E, Wp = expected(P, i, ch["method"], ch["cfg"], cache, tg_g)
    if ch["level"] == "total" or E.shape[0] == 1:
        cums = _paths(A.sum(axis=0)[None], E.sum(axis=0)[None], ch["nc"])
    else:
        cums = _paths(A, E, ch["nc"])
    w1 = _w1_paths(W, Wp) if Wp.shape[0] == W.shape[0] else _w1_paths(W.sum(axis=0)[None], Wp)
    return cums, w1, E


def path_rows(A_tot, W_tot, TG, cums, w1) -> list:
    """Разница прогноза и факта при k = 0..n-1: выполнение плана и RR на конец месяца и на каждую дату после k, п.п."""
    n = len(A_tot)
    day = np.arange(1, n + 1)
    with np.errstate(divide="ignore", invalid="ignore"):
        comp_f, rr_f = A_tot / TG, A_tot / W_tot
        d_comp = 100 * (cums / TG - comp_f)
        d_rr = 100 * (cums / w1 - rr_f)
    rows = []
    for k in range(n):
        dc, dr = d_comp[k, k:], d_rr[k, k:]
        ok, ok5 = np.isfinite(dc), np.isfinite(dc) & (day[k:] >= 5)
        rows.append(dict(known_days=k, comp_fact=100 * comp_f[-1], rr_fact=100 * rr_f[-1],
                         err_comp=d_comp[k, -1], err_rr=d_rr[k, -1],
                         dates_comp=np.abs(dc[ok]).mean() if ok.any() else np.nan,
                         dates_comp_5=np.abs(dc[ok5]).mean() if ok5.any() else np.nan,
                         dates_rr=np.abs(dr[np.isfinite(dr)]).mean() if np.isfinite(dr).any() else np.nan,
                         dates_within1=(np.abs(dc[ok]) <= 1).mean() if ok.any() else np.nan))
    return rows


def honest(P: dict, test: list, scores: list, cache: dict) -> dict:
    """Честный бэктест: месяц test[t] прогнозируется вариантами, выбранными по test[:t]."""
    rows, paths, picks = [], {}, {}
    for t in range(min(max(HONEST_WARMUP, 1), len(test) - 1), len(test)):
        best, winner = select(P, test, scores, list(range(t)), cache)
        i = test[t]
        A, W, TG = _actual(P, i)
        for name in list(best) + ["auto"]:
            ch = best[winner if name == "auto" else name]
            cums, w1, _ = forecast_paths(P, i, ch, cache, A, W, plan_known(P, i, cache))
            for r in path_rows(A.sum(axis=0), W.sum(axis=0), TG, cums, w1):
                rows.append(dict(method=name, ym=str(P["months"][i]), **r))
            if name == "auto":
                paths[str(P["months"][i])], picks[str(P["months"][i])] = (cums, w1), winner
    return dict(df=pd.DataFrame(rows), paths=paths, picks=picks)


# ============================================================================
# ПРОФИЛЬ ДАННЫХ
# ============================================================================
def _mix_rate(P: dict) -> pd.DataFrame:
    """Разложение изменения RR конца месяца от месяца к месяцу: структура (доли w1 сегментов) и сбор внутри сегментов."""
    rows = []
    prev = None
    for i in np.flatnonzero(P["complete"]):
        n = P["dim"][i]
        w, c = P["W1"][i, :, n - 1], P["C"][i, :, n - 1]
        s = w / w.sum()
        with np.errstate(divide="ignore", invalid="ignore"):
            r = np.where(w > 0, c / np.where(w > 0, w, 1), 0.0)
        if prev is not None:
            ps, pr = prev
            rows.append(dict(ym=str(P["months"][i]), rr=100 * (s * r).sum(), d_rr=100 * ((s * r).sum() - (ps * pr).sum()),
                             mix=100 * ((s - ps) * pr).sum(), rate=100 * (s * (r - pr)).sum()))
        prev = (s, r)
    return pd.DataFrame(rows).set_index("ym")


def profile(raw: pd.DataFrame, P: dict = None) -> dict:
    """Как устроены данные: печатает диагностику, по которой настраивается модель. Пришлите вывод целиком."""
    P = P or build_panel(raw)
    c = COLS
    months = P["months"]
    print(f"Строк: {len(raw):,}".replace(",", " ") + (f"; факт (is_forecast = 0): {(raw[c['is_forecast']] == 0).sum():,}, "
          f"план на будущее (= 1): {(raw[c['is_forecast']] == 1).sum():,}".replace(",", " ") if c["is_forecast"] in raw else ""))
    print(f"Месяцы: {months[0]} – {months[-1]}; полных месяцев с фактом: {P['complete'].sum()}; "
          f"текущий месяц: {months[~P['complete'] & P['has'].any(axis=1)].astype(str).tolist() or '—'}")
    print(f"Сегментов прогноза (GROUP_COLS = {GROUP_COLS}): {len(P['groups'])}; уровни подтягивания: "
          + " → ".join(f"{lev['L']} ({' × '.join(lev['cols']) or 'портфель'})" for lev in P["levels"]))

    d = pd.to_datetime(raw.loc[raw[c["is_forecast"]] == 0, c["date"]] if c["is_forecast"] in raw else raw[c["date"]])
    ud = pd.Series(d.unique())
    print("\n[1] Срезы: дней со срезом в месяце (мин / медиана / макс): "
          f"{P['has'].sum(axis=1)[P['complete']].min()} / {int(np.median(P['has'].sum(axis=1)[P['complete']]))} / "
          f"{P['has'].sum(axis=1)[P['complete']].max()}; доля выходных среди дат срезов: "
          f"{100 * (ud.dt.day_of_week >= 5).mean():.0f}%; первый срез месяца — обычно {int(np.median(P['first'][P['complete']])) + 1}-е число, "
          f"последний — {int(np.median(P['last'][P['complete']])) + 1}-е")
    if "dt_in" in raw and "inflow" in raw:
        same = pd.to_datetime(raw["dt_in"]).dt.to_period("M") == pd.to_datetime(raw[c["month"]]).dt.to_period("M")
        t = same.groupby(raw["inflow"]).mean().mul(100).round(1)
        print("[2] dt_in в том же месяце, что report_month, % строк — по inflow: " + ", ".join(f"{k}: {v}%" for k, v in t.items()))
    vary = P["varying"]
    g = P["groups"]
    msg = f"[3] w1 растёт внутри месяца у {vary.sum()} из {len(vary)} сегментов"
    if "inflow" in g:
        msg += " (по inflow: " + ", ".join(f"{k}: {v:.0f}%" for k, v in
                                           pd.Series(vary).groupby(g["inflow"].to_numpy()).mean().mul(100).items()) + " сегментов)"
    ws = np.array([P["W1"][i, :, P["dim"][i] - 1].sum() for i in np.flatnonzero(P["complete"])])
    wv = np.array([P["W1"][i, vary, P["dim"][i] - 1].sum() for i in np.flatnonzero(P["complete"])])
    print(msg + f"; их доля в w1 конца месяца: {100 * np.median(wv / ws):.1f}%")
    first_rr = []
    for i in np.flatnonzero(P["complete"]):
        f = P["first"][i]
        first_rr.append(100 * P["C"][i, :, f].sum() / max(P["W1"][i, :, f].sum(), 1e-9))
    print(f"[4] RR на первом срезе месяца (медиана): {np.median(first_rr):.2f}% — если около 0, когорта месячная "
          f"(сбор считается с начала месяца)")

    rows = []
    for i in np.flatnonzero(P["has"].any(axis=1)):
        n = P["dim"][i]
        q = n - 1 if P["complete"][i] else P["last"][i]
        w, cc, tg = P["W1"][i, :, q].sum(), P["C"][i, :, q].sum(), P["TG"][i, :, q].sum()
        rows.append({"месяц": str(months[i]), "по число": q + 1, "w1 первый срез": P["W1"][i, :, P["first"][i]].sum() / 1e9,
                     "w1 на дату": w / 1e9, "сбор на дату": cc / 1e9, "RR, %": 100 * cc / w if w else np.nan,
                     "target на дату": tg / 1e9, "целевой RR, %": 100 * tg / w if w else np.nan,
                     "выполнение, %": 100 * cc / tg if tg else np.nan})
    tab = pd.DataFrame(rows).set_index("месяц")
    print("\n[5] По месяцам (суммы — млрд руб.; полные месяцы — на конец, текущий — на последний срез):")
    print(tab.round(2).to_string())

    if "dpd_bucket" in GROUP_COLS:
        last = np.flatnonzero(P["complete"])[-6:]
        k = g["dpd_bucket"].to_numpy()
        rr_b = {}
        for b in sorted(set(k)):
            m = k == b
            w = sum(P["W1"][i, m, P["dim"][i] - 1].sum() for i in last)
            cc = sum(P["C"][i, m, P["dim"][i] - 1].sum() for i in last)
            rr_b[b] = {"доля w1, %": 100 * w / sum(P["W1"][i, :, P["dim"][i] - 1].sum() for i in last),
                       "RR за месяц, %": 100 * cc / w if w else np.nan}
        print("\n[6] DPD-бакеты, последние 6 полных месяцев:")
        print(pd.DataFrame(rr_b).T.round(2).to_string())
    mr = _mix_rate(P)
    if len(mr):
        print(f"\n[7] Изменение RR конца месяца от месяца к месяцу, средний модуль, п.п.: всего {mr['d_rr'].abs().mean():.2f}; "
              f"из-за структуры портфеля {mr['mix'].abs().mean():.2f}; из-за сбора внутри сегментов {mr['rate'].abs().mean():.2f}.")
        print("    Структуру прогноз снизу вверх учитывает по известному w1, прогноз по итогу — нет.")
    if P["fhas"].any():
        fm = [str(months[i]) for i in np.flatnonzero(P["fhas"].any(axis=1))]
        print(f"[8] Строки плана на будущее (is_forecast = 1) есть за: {', '.join(fm)}")
    else:
        print("[8] Строк плана на будущее (is_forecast = 1) нет: план будущих дней текущего месяца будет продолжен "
              "равномерно по дням")
    return dict(panel=P, months=tab, mix_rate=mr)


# ============================================================================
# РАБОЧИЙ ПРОГНОЗ
# ============================================================================
def _plan_path(P: dict, i: int, k: int) -> tuple:
    """План сегментов [G, n] на каждый день текущего месяца: факт по сегодня, дальше — строки плана на будущее
    (is_forecast = 1), а если их нет — продолжен равномерно по дням."""
    n = P["dim"][i]
    tg = np.nan_to_num(P["TG"][i, :, :n])
    day = np.arange(n)
    if P["fhas"][i].any() and "target" in P["fut"]:
        return np.where(day < k, tg, P["fut"]["target"][i, :, :n]), "строки плана на будущее"
    if k > 0:
        return np.where(day < k, tg, tg[:, [k - 1]] * (day + 1) / k), "продолжен равномерно по дням"
    return np.full((len(tg), n), np.nan), "нет"


def production(P: dict, best: dict, winner: str, cache: dict, hon: dict) -> tuple:
    """Прогноз текущего месяца по каждому дню: факт по сегодня + прогноз остатка; плюс прогноз с 1-го числа."""
    i = int(np.flatnonzero(~P["complete"] & P["has"].any(axis=1))[0]) if (~P["complete"] & P["has"].any(axis=1)).any() else None
    if i is None:
        print("Текущего (неполного) месяца в данных нет — рабочего прогноза нет, только бэктест.")
        return None, None, None
    n, k = P["dim"][i], P["last"][i] + 1
    ch = best[winner]
    tg_g, how = _plan_path(P, i, k)
    TG = tg_g.sum(axis=0)
    A = np.nan_to_num(P["C"][i, :, :n])
    W = np.nan_to_num(P["W1"][i, :, :n])
    cums, w1, E = forecast_paths(P, i, ch, cache, A, W, tg_g)
    dates = pd.date_range(P["months"][i].start_time, periods=n, freq="D")
    has = np.arange(n) < k
    a_tot, w_tot = A.sum(axis=0), W.sum(axis=0)
    o = pd.DataFrame({"month": str(P["months"][i]), "date": dates, "has_fact": has,
                      "target": TG, "w1_fact": np.where(has, w_tot, np.nan), "w1": w1[min(k, n - 1)],
                      "collected_fact": np.where(has, a_tot, np.nan),
                      "collected_forecast": cums[0], "collected_fact_plus_forecast": cums[min(k, n - 1)]})
    with np.errstate(divide="ignore", invalid="ignore"):
        o["comp_fact_pct"] = 100 * o["collected_fact"] / o["target"]
        o["comp_forecast_pct"] = 100 * o["collected_forecast"] / o["target"]
        o["comp_fact_plus_forecast_pct"] = 100 * o["collected_fact_plus_forecast"] / o["target"]
        o["rr_fact_pct"] = 100 * o["collected_fact"] / o["w1_fact"]
        o["rr_forecast_pct"] = 100 * o["collected_forecast"] / w1[0]
        o["rr_fact_plus_forecast_pct"] = 100 * o["collected_fact_plus_forecast"] / o["w1"]
    # интервал из честного бэктеста «Авто» при k известных днях
    h = hon["df"]
    a = h[(h["method"] == "auto") & (h["known_days"] == min(k, h["known_days"].max()))]
    comp_end = o["comp_fact_plus_forecast_pct"].iloc[-1]
    rr_end = o["rr_fact_plus_forecast_pct"].iloc[-1]
    q = (lambda s, x: (x - s.quantile(0.9), x - s.quantile(0.1))) if len(a) >= 5 else (lambda s, x: (np.nan, np.nan))
    lo, hi = q(a["err_comp"], comp_end)
    lo_r, hi_r = q(a["err_rr"], rr_end)
    summ = {"месяц": str(P["months"][i]), "срезов": f"по {k}-е из {n}", "план дальше": how,
            "выполнение: прогноз с 1-го": f"{o['comp_forecast_pct'].iloc[-1]:.1f}%",
            "выполнение: факт + прогноз": f"{comp_end:.1f}%",
            "80% интервал выполнения": f"{lo:.1f}% – {hi:.1f}%" if np.isfinite(lo) else "—",
            "RR: факт + прогноз": f"{rr_end:.2f}%",
            "80% интервал RR": f"{lo_r:.2f}% – {hi_r:.2f}%" if np.isfinite(lo_r) else "—",
            "метод": NAMES[winner] + f" ({describe(ch)})"}
    seg = None
    if E.shape[0] > 1:     # сегменты: где факт отстаёт от ожидания
        Eg = E
        q_ = max(k - 1, 0)
        seg = P["groups"].copy()
        seg["w1"] = W[:, q_]
        seg["сбор на сегодня"] = A[:, q_]
        seg["ожидание на сегодня"] = Eg[:, q_]
        seg["отклонение, руб"] = seg["сбор на сегодня"] - seg["ожидание на сегодня"]
        seg["RR на сегодня, %"] = 100 * seg["сбор на сегодня"] / seg["w1"].replace(0, np.nan)
        seg["ожидаемый RR на конец, %"] = 100 * Eg[:, -1] / seg["w1"].replace(0, np.nan)
        seg = seg.sort_values("отклонение, руб")
    return o, summ, seg


def describe(ch: dict) -> str:
    a, tau = ch["nc"]
    corr = "без поправки" if a == 0 else f"поправка {'по сегментам' if ch['level'] == 'group' else 'по портфелю'}: α_max {a:g}, τ {tau:g}"
    if ch["method"] == "ref":
        return "3 месяца по числу месяца; без поправки"
    if ch["method"] == "mix":
        w, cbu, cbp = ch["cfg"]
        return (f"{w:g} · от сбора ({_cfg_text(cbu, True)}) + {1 - w:g} · от плана ({_cfg_text(cbp, True, False)}); "
                + corr)
    return _cfg_text(ch["cfg"], ch["method"] in ("bu", "bp"), ch["method"] != "bp") + "; " + corr


def _cfg_text(cfg, shrink: bool, seasonal: bool = True) -> str:
    how, h, K, s = cfg
    parts = [f"дни {'по календарю' if how == 'calendar' else 'по рабочим дням'}",
             "месяцы поровну" if np.isinf(h) else f"полураспад {h:g} мес."]
    if seasonal:
        parts.append(f"сезон {s:g}")
    if shrink:
        parts.append(f"подтягивание K {K:g}")
    return ", ".join(parts)


# ============================================================================
# ОТЧЁТ И EXCEL
# ============================================================================
def honest_table(h: pd.DataFrame) -> pd.DataFrame:
    res = {}
    for name in list(METHOD_CFGS) + ["mix", "auto"]:
        g = h[h["method"] == name]
        if g.empty:
            continue
        by_k = g.groupby("known_days")
        mc = by_k["err_comp"].apply(lambda s: s.abs().mean())
        mr = by_k["err_rr"].apply(lambda s: s.abs().mean())
        within = by_k["err_comp"].apply(lambda s: (s.abs() <= 1).mean())
        g0 = g[g["known_days"] == 0]
        res[NAMES[name]] = {
            **{f"выполнение, конец мес., известно {k} дн.": f"{mc.get(k, np.nan):.2f} п.п." for k in REPORT_KNOWN_DAYS},
            **{f"RR, конец мес., известно {k} дн.": f"{mr.get(k, np.nan):.2f} п.п." for k in (0, 10, 20)},
            "месяцев с выполнением ≤ 1 п.п. (0 / 10 / 20 дн.)": " / ".join(f"{100 * within.get(k, np.nan):.0f}%" for k in (0, 10, 20)),
            "на даты с 5-го числа, прогноз с 1-го: выполнение": f"{g0['dates_comp_5'].mean():.2f} п.п.",
            "на даты, прогноз с 1-го: RR": f"{g0['dates_rr'].mean():.2f} п.п.",
        }
    return pd.DataFrame(res)


def _triangles(P: dict, fc: pd.DataFrame, hon: dict):
    """Треугольники месяц × число месяца: выполнение и RR (факт; в текущем месяце дальше — прогноз), разница прогноза
    с 1-го числа и факта по дням (честный бэктест), прогноз итога по мере факта."""
    comp, rr, is_fc = {}, {}, {}
    for i in np.flatnonzero(P["has"].any(axis=1)):
        n = P["dim"][i]
        m = str(P["months"][i])
        q = n if P["complete"][i] else P["last"][i] + 1
        c, w, t = P["C"][i, :, :n].sum(axis=0), P["W1"][i, :, :n].sum(axis=0), P["TG"][i, :, :n].sum(axis=0)
        with np.errstate(divide="ignore", invalid="ignore"):
            comp[m] = np.where(np.arange(n) < q, c / t, np.nan)
            rr[m] = np.where(np.arange(n) < q, c / w, np.nan)
        is_fc[m] = np.zeros(n, bool)
    if fc is not None:
        m = fc["month"].iloc[0]
        comp[m] = (fc["comp_fact_plus_forecast_pct"] / 100).to_numpy()
        rr[m] = (fc["rr_fact_plus_forecast_pct"] / 100).to_numpy()
        is_fc[m] = ~fc["has_fact"].to_numpy()

    def tri(d):
        t = pd.DataFrame({m: pd.Series(v, index=range(1, len(v) + 1)) for m, v in d.items()}).T
        return t.reindex(columns=range(1, 32))

    diff_c, diff_r, conv = {}, {}, {}
    for m, (cums, w1) in hon["paths"].items():
        i = list(P["months"].astype(str)).index(m)
        A, W, TG = _actual(P, i)
        a, w = A.sum(axis=0), W.sum(axis=0)
        with np.errstate(divide="ignore", invalid="ignore"):
            diff_c[m] = 100 * (cums[0] - a) / TG
            diff_r[m] = 100 * (cums[0] / w1[0] - a / w)
            conv[m] = pd.Series(cums[:, -1] / TG[-1])
    return tri(comp), tri(rr), tri(is_fc).astype(float).fillna(0).astype(bool), tri(diff_c), tri(diff_r), conv


def export_excel(path: str, P: dict, fc, summ, seg, hon: dict, tables: list):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("Excel не записан: нужен пакет openpyxl (pip install openpyxl).")
        return
    st = dict(fc_fill=PatternFill("solid", fgColor="DCEBFA"), fc_font=Font(color="1F5FAD", italic=True),
              none=PatternFill("solid", fgColor="F2F2F2"), ok=PatternFill("solid", fgColor="D9F2E3"),
              mid=PatternFill("solid", fgColor="FFF2CC"), bad=PatternFill("solid", fgColor="F8D7D3"))
    comp, rr, isfc, dc, dr, conv = _triangles(P, fc, hon)

    def dim(m):
        return pd.Period(m, "M").days_in_month if str(m)[:1].isdigit() else 31

    def is_day(c):
        return isinstance(c, (int, np.integer))

    def style_tri(mask):
        def s(idx, col, v):
            if is_day(col):
                if col > dim(idx):
                    return st["none"], None, None
                if idx in mask.index and mask.loc[idx, col]:
                    return st["fc_fill"], st["fc_font"], None
            return None, None, None
        return s

    def style_err(idx, col, v):
        if is_day(col) and col > dim(idx):
            return st["none"], None, None
        if idx == "Месяцев ≤ 1 п.п.":
            return None, Font(bold=True), cf.XL_PCT
        return cf._err_fill(st, v), (Font(bold=True) if idx == "Средняя |разница|" else None), None

    wb = Workbook()
    sheets = [("Выполнение плана", comp, "Выполнение плана на дату: сбор с начала месяца / target на дату", style_tri(isfc), cf.XL_PCT),
              ("RR", rr, "RR на дату: сбор с начала месяца / w1", style_tri(isfc), "0.00%")]
    for k, (name, t, title, style, fmt) in enumerate(sheets):
        ws = wb.active if k == 0 else wb.create_sheet(name)
        ws.title = name
        t = t.copy()
        t.index.name = "Месяц"
        cf._put_table(ws, t, 1, title=title, note="Чёрным — факт, синим на голубом — прогноз (факт по сегодня + прогноз "
                      "остатка). Столбцы 1–31 — число месяца.", fmt=lambda c, f=fmt: f, style=style)
    for name, t, title in (("Разница выполнения", dc, "Прогноз с 1-го числа минус факт: выполнение плана на дату, п.п."),
                           ("Разница RR", dr, "Прогноз с 1-го числа минус факт: RR на дату, п.п.")):
        if t.empty:
            continue
        t = t.copy()
        base = t.copy()
        t.loc["Средняя |разница|"] = base.abs().mean()
        t.loc["Месяцев ≤ 1 п.п."] = (base.abs() <= 1).sum() / base.notna().sum().replace(0, np.nan)
        t.index.name = "Месяц"
        cf._put_table(wb.create_sheet(name), t, 1, title=title,
                      note="Честный бэктест «Авто». Зелёный — в пределах 1 п.п., жёлтый — до 2, красный — больше.",
                      fmt=lambda c: cf.XL_PP, style=style_err)
    if conv:
        K = pd.DataFrame(conv).T.reindex(columns=range(31))
        fact = {m: comp.loc[m].dropna().iloc[-1] if m in comp.index else np.nan for m in K.index}
        K["Факт на конец"] = [fact[m] for m in K.index]
        K.index.name = "Месяц"

        def s4(idx, col, v):
            if col == "Факт на конец":
                return None, Font(bold=True), None
            if col >= dim(idx):
                return st["none"], None, None
            f = K.loc[idx, "Факт на конец"]
            return cf._err_fill(st, 100 * (v - f) if v is not None and np.isfinite(f) else None), None, None

        cf._put_table(wb.create_sheet("Итог по мере факта"), K, 1,
                      title="Прогноз выполнения плана на конец месяца, когда известно k дней месяца (честный бэктест)",
                      note="Цвет — разница с фактом на конец месяца: зелёный ≤ 1 п.п., жёлтый ≤ 2, красный больше.",
                      fmt=lambda c: cf.XL_PCT, style=s4)
    if seg is not None:
        cf._put_table(wb.create_sheet("Сегменты сейчас"), seg.reset_index(drop=True), 1, index=False,
                      title="Текущий месяц по сегментам: факт против ожидания на сегодня (сверху — сильнее всего отстают)",
                      fmt=lambda c: "0.00" if "%" in str(c) else (cf.XL_MONEY if c in ("w1", "сбор на сегодня",
                                                                                      "ожидание на сегодня", "отклонение, руб") else None))
    for name, parts in tables:
        ws = wb.create_sheet(name[:31])
        r = 1
        for title, t in parts:
            r = cf._put_table(ws, t, r, title=title, index="date" not in t.columns,
                              fmt=lambda c: "0.00" if any(s in str(c) for s in ("%", "п.п.", "pct")) else
                              (cf.XL_MONEY if str(c) in ("target", "w1", "w1_fact") or str(c).startswith("collected") else None))
    for ws in wb.worksheets:
        ws.column_dimensions["A"].width = 12
        for j in range(2, ws.max_column + 1):
            head = ws.cell(3, j).value
            ws.column_dimensions[get_column_letter(j)].width = 8 if str(head).isdigit() else 15
        if ws.title in ("Выполнение плана", "RR", "Разница выполнения", "Разница RR", "Итог по мере факта"):
            ws.freeze_panes = "B4"
    wb.save(path)
    print(f"Excel: {path}")


# ============================================================================
# ЗАПУСК
# ============================================================================
def run(raw: pd.DataFrame, P: dict = None) -> dict:
    P = P or build_panel(raw)
    months = P["months"]
    start = pd.Period(BACKTEST_FROM, "M")
    test = [i for i in np.flatnonzero(P["complete"]) if months[i] >= start
            and sum(P["complete"][max(0, i - HISTORY_MONTHS):i]) >= 3]
    if len(test) < 3:
        raise ValueError("Мало полных месяцев для бэктеста — сдвиньте BACKTEST_FROM раньше")
    print(f"Сегментов: {len(P['groups'])}; бэктест: {months[test[0]]} – {months[test[-1]]} ({len(test)} мес.)")
    cache = {}
    scores = []
    for i in test:
        scores.append(score_month(P, i, cache))
        print(f"  бэктест {months[i]}")
    hon = honest(P, test, scores, cache)
    best, winner = select(P, test, scores, list(range(len(test))), cache)
    h = hon["df"]
    hm = sorted(h["ym"].unique())
    print(f"\nЧЕСТНЫЙ БЭКТЕСТ: каждый месяц ({len(hm)} мес., {hm[0]} – {hm[-1]}) прогнозируется методом и параметрами, "
          f"выбранными только по предыдущим месяцам. Средний модуль разницы прогноза и факта, п.п.")
    ht = honest_table(h)
    print(ht.to_string())
    a = h[h["method"] == "auto"]
    by_month = a[a["known_days"].isin(REPORT_KNOWN_DAYS)].pivot(index="ym", columns="known_days", values="err_comp").round(2)
    by_month.insert(0, "выполнение факт, %", a.groupby("ym")["comp_fact"].first().round(1))
    by_month.insert(0, "метод", pd.Series(hon["picks"]).map(NAMES))
    print("\n«Авто» по месяцам: разница выполнения плана на конец месяца, п.п. (столбцы — сколько дней известно; "
          "+ — прогноз выше факта):")
    print(by_month.to_string())
    print(f"\nРабочий прогноз — параметры подобраны по всем {len(test)} мес. бэктеста:")
    print(pd.DataFrame({NAMES[m]: {"параметры": describe(b), "ошибка на подборе": f"{b['score']:.2f} п.п."}
                        for m, b in best.items()}).T.to_string())
    print(f"→ в прогнозе: {NAMES[winner]}")
    fc, summ, seg = production(P, best, winner, cache, hon)
    if summ:
        print("\nПрогноз текущего месяца:")
        print(pd.Series(summ).to_string())
        if seg is not None:
            print("\nСегменты, сильнее всего отстающие от ожидания на сегодня:")
            print(seg.head(10).round(2).to_string(index=False))
    tables = [("Точность методов", [("Честный бэктест: средний модуль разницы прогноза и факта, п.п.", ht)]),
              ("Авто по месяцам", [("«Авто» по месяцам: разница выполнения плана на конец месяца, п.п.",
                                    by_month.rename_axis("Месяц").rename(columns={k: f"известно {k} дн., п.п." for k in REPORT_KNOWN_DAYS}))])]
    if summ:
        tables = [("Прогноз", [("Прогноз текущего месяца", pd.Series(summ, name="значение").to_frame())])] + tables
        tables.append(("Прогноз по дням", [("Текущий месяц по дням", fc)]))
    prof = profile_tables(P)
    tables.append(("Профиль данных", prof))
    if EXCEL:
        export_excel(f"{OUT_DIR}/rr_effectiveness.xlsx", P, fc, summ, seg, hon, tables)
    return dict(panel=P, scores=scores, test=test, honest=hon, best=best, winner=winner, forecast=fc, summary=summ,
                segments=seg, table=ht, by_month=by_month)


def profile_tables(P: dict) -> list:
    """Таблицы профиля для Excel (без печати)."""
    rows = []
    for i in np.flatnonzero(P["has"].any(axis=1)):
        n = P["dim"][i]
        q = n - 1 if P["complete"][i] else P["last"][i]
        w, cc, tg = P["W1"][i, :, q].sum(), P["C"][i, :, q].sum(), P["TG"][i, :, q].sum()
        rows.append({"месяц": str(P["months"][i]), "по число": q + 1, "w1, млрд": w / 1e9, "сбор, млрд": cc / 1e9,
                     "RR, %": 100 * cc / w if w else np.nan, "выполнение, %": 100 * cc / tg if tg else np.nan})
    out = [("Месяцы", pd.DataFrame(rows).set_index("месяц").round(3))]
    mr = _mix_rate(P)
    if len(mr):
        out.append(("Изменение RR от месяца к месяцу, п.п.: структура портфеля и сбор внутри сегментов",
                    mr.round(3).rename(columns={"rr": "RR, %", "d_rr": "изменение", "mix": "структура", "rate": "сбор в сегментах"})))
    return out


if __name__ == "__main__":
    raw = load_raw()
    profile(raw)
    run(raw)
